"""Synthetic approved allocations remain distinct in frozen files and archives."""

import json
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4
from zipfile import ZipFile

import pytest
from conftest import API
from factoryflow.bootstrap import create_app
from factoryflow.kernel.contracts import BusinessError
from factoryflow.modules.statements.adapters.archive_repository import ArchiveStatementRepository
from factoryflow.modules.statements.adapters.excel import render_excel
from factoryflow.modules.statements.adapters.export_repository import StatementExportRepository
from factoryflow.modules.statements.application.archive import ArchiveStatements
from openpyxl import load_workbook
from sqlalchemy import text
from test_postgres_statement_lifecycle import (
    approve,
    blank,
    command,
    detail,
    reserve,
    settle,
    setup_statement,
)
from test_postgres_statements import draft, source
from test_statement_excel import snapshot


def allocated_snapshot(status="confirmed"):
    data = snapshot()
    statement = data["statement"]
    statement["status"] = status
    ordinary = statement["lines"][0]
    ordinary.update(status="released", quantity="0", unit_price="10")
    statement["settlement_lines"] = [
        {
            **ordinary,
            "status": "included",
            "quantity": quantity,
            "unit_price": price,
            "allocation_id": f"allocation-{index}",
            "approval_id": f"approval-{index}",
            "approval_version": 1,
            "allocation_status": "settled" if status == "confirmed" else "reserved",
        }
        for index, (quantity, price) in enumerate([("60", "10"), ("30", "12")], 1)
    ]
    statement["calculation"] = {
        "line_amounts": ["0.00", "600.00", "360.00"],
        "priced_subtotal": "960.00",
        "tax_base": "960.00",
        "tax_amount": "0.00",
        "total": "960.00",
        "warnings": [],
    }
    return data


def test_formal_excel_contains_separate_allocations_and_exact_calculation_indices():
    book = load_workbook(BytesIO(render_excel(allocated_snapshot())))
    sheet = book.active
    assert "正式" in sheet["A1"].value
    assert sheet["G7"].value == 60 and sheet["G8"].value == 30
    assert sheet["H7"].value == 10 and sheet["H8"].value == 12
    assert sheet["I7"].value == 600 and sheet["I8"].value == 360
    assert "approval-1" in sheet["J7"].value
    assert "allocation-2" in sheet["J8"].value
    assert not any(cell.data_type == "f" for row in sheet for cell in row)


def test_grouped_historical_excel_keeps_calculation_mapping_after_sort():
    data = allocated_snapshot("draft")
    data.update(template_id="delivery-v1", current=False)
    data["statement"]["settlement_lines"][0]["source"] = {
        **data["statement"]["lines"][0]["source"],
        "delivery_date": "2026-10-02",
    }
    data["statement"]["settlement_lines"][0]["delivery_id"] = "later-delivery"
    sheet = load_workbook(BytesIO(render_excel(data))).active
    assert "历史版本" in sheet["A1"].value
    rows = [row for row in sheet if row[6].value in (30, 60)]
    assert [(row[6].value, row[8].value) for row in rows] == [(30, 360), (60, 600)]
    data["statement"]["calculation"]["line_amounts"] = ["600.00"]
    with pytest.raises(ValueError, match="calculation"):
        render_excel(data)


def test_formal_export_and_archive_keep_fixed_bytes_after_new_formal_version(pg):
    delivery = source(pg)
    created = command(pg, "/api/v1/statements", draft(pg, delivery))
    sent_export = command(
        pg, f"/api/v1/statements/{created['id']}/exports", {"expected_version": 1}
    )
    assert pg.runtime.statement_exports.generate(pg.actor, sent_export["id"])["status"] == "ready"
    sent = command(
        pg,
        f"/api/v1/statements/{created['id']}/sent",
        {
            "expected_version": 1,
            "export_id": sent_export["id"],
            "person": "Synthetic Sender",
            "occurred_on": "2026-10-02",
            "note": "Synthetic real sent export",
        },
    )
    feedback = command(
        pg,
        f"/api/v1/statements/{created['id']}/feedback",
        {
            "person": "Synthetic Contact",
            "occurred_on": "2026-10-03",
            "note": "Synthetic phone feedback",
        },
    )
    params = {
        "org": pg.actor.organization_id,
        "actor": pg.actor.user_id,
        "customer": pg.customer_id,
        "statement": created["id"],
        "delivery": delivery["id"],
        "line": delivery["lines"][0]["order_line_id"],
        "sent": sent["id"],
        "feedback": feedback["id"],
        "approval": uuid4(),
        "source": json.dumps(created["lines"][0]["source"]),
    }
    with pg.owner.begin() as conn:
        conn.execute(
            text("""
            UPDATE statement_lines SET quantity=0,status='released',reason='Synthetic approval'
            WHERE organization_id=:org AND statement_id=:statement
        """),
            params,
        )
        conn.execute(
            text("""
            INSERT INTO statement_approvals(organization_id,id,customer_id,statement_id,
            delivery_id,order_line_id,sent_id,feedback_id,quantity,unit_price,target_start,
            target_end,reason,source,actor_id) VALUES(:org,:approval,:customer,:statement,
            :delivery,:line,:sent,:feedback,0.1,2,'2026-10-01','2026-10-31',
            'Synthetic direct database approval',CAST(:source AS jsonb),:actor)
        """),
            params,
        )
    data = pg.runtime.statements.detail(pg.actor, UUID(params["statement"]))
    data.update(status="confirmed", version=2, fees_confirmed=True, reconfirmation="not_required")
    allocation = uuid4()
    data["settlement_lines"] = [
        {
            **data["lines"][0],
            "status": "included",
            "quantity": "0.06",
            "unit_price": "2",
            "allocation_id": str(allocation),
            "approval_id": str(params["approval"]),
            "approval_version": 1,
            "allocation_status": "settled",
        }
    ]
    data["calculation"] = {
        "line_amounts": ["0.00", "0.12"],
        "priced_subtotal": "0.12",
        "tax_base": "0.12",
        "tax_amount": "0.00",
        "total": "0.12",
        "warnings": [],
    }
    identifier = uuid4()
    frozen = {
        "id": str(identifier),
        "template_id": "standard-v1",
        "template_version": 1,
        "calculation_version": 1,
        "statement": data,
    }
    with pg.owner.begin() as conn:
        conn.execute(
            text("""
            UPDATE statements SET status='confirmed',version=2,fees_confirmed=true,
            reconfirmation='not_required' WHERE id=:statement
        """),
            params,
        )
        conn.execute(
            text("""
            INSERT INTO statement_events(organization_id,id,statement_id,version,action,
            reason,actor_id,snapshot) VALUES(:org,:event,:statement,2,'settle',
            'Synthetic formal output',:actor,CAST(:snapshot AS jsonb))
        """),
            params | {"event": uuid4(), "snapshot": json.dumps(data)},
        )
        conn.execute(
            text("""
            INSERT INTO statement_allocations(organization_id,id,customer_id,approval_id,
            statement_id,quantity,unit_price,status,settlement_version,reason,actor_id)
            VALUES(:org,:allocation,:customer,:approval,:statement,0.06,2,'settled',2,
            'Synthetic formal allocation',:actor)
        """),
            params | {"allocation": allocation},
        )
        StatementExportRepository(conn).create(pg.actor, frozen)
    service = pg.runtime.statement_exports
    ready = service.generate(pg.actor, identifier)
    assert ready["status"] == "ready" and ready["current"]
    assert ready["statement_status"] == "confirmed"
    content, _ = service.download(pg.actor, identifier)
    assert "正式" in load_workbook(BytesIO(content)).active["A1"].value
    with pg.owner.connect() as conn:
        facade = ArchiveStatements(ArchiveStatementRepository(conn), service.store)
        metadata, archived = facade.capture(pg.actor, f"statement_export:{identifier}")
        assert metadata["current"] and "正式" in metadata["title"]
        package = archived
        with ZipFile(BytesIO(package)) as zipped:
            assert zipped.read(f"statements/{data['id']}/v2/statement.xlsx") == content
    with pg.owner.begin() as conn:
        conn.execute(text("UPDATE statements SET version=3 WHERE id=:statement"), params)
        conn.execute(
            text("""
            INSERT INTO statement_events(organization_id,id,statement_id,version,action,
            reason,actor_id,snapshot) VALUES(:org,:event,:statement,3,'correct_settlement',
            'Synthetic new formal version',:actor,CAST(:snapshot AS jsonb))
        """),
            params | {"event": uuid4(), "snapshot": json.dumps(data | {"version": 3})},
        )
    historical = service.generate(pg.actor, identifier)
    assert historical["file_hash"] == ready["file_hash"] and not historical["current"]
    with pytest.raises(BusinessError, match="STATEMENT_EXPORT_STALE"):
        service.download(pg.actor, identifier)
    assert service.download(pg.actor, identifier, historical=True)[0] == content
    with pg.owner.connect() as conn:
        facade = ArchiveStatements(ArchiveStatementRepository(conn), service.store)
        metadata, archived = facade.capture(pg.actor, f"statement_export:{identifier}")
        assert not metadata["current"] and archived == package


def test_http_formal_export_archive_and_later_review_keep_original_bytes(pg, tmp_path):
    application = create_app(settings=replace(pg.settings, archive_root=tmp_path / "copies"))
    api = API(application)
    api.login()
    flow = SimpleNamespace(api=api, customer_id=pg.customer_id)
    try:
        statement, source, sent, feedback = setup_statement(flow)
        approval = approve(flow, statement, source, sent, feedback)
        review = command(
            flow,
            f"/api/v1/statements/{statement['id']}/review",
            {
                "expected_version": detail(flow, statement["id"])["version"],
                **source,
                "quantity": "10",
                "reason": "Synthetic disputed quantity",
            },
        )
        reserve(flow, statement, approval, "60")
        formal = settle(flow, statement)
        export = command(
            flow,
            f"/api/v1/statements/{statement['id']}/exports",
            {
                "expected_version": formal["version"],
            },
        )
        ready = api.request("POST", f"/api/v1/statement-exports/{export['id']}/generate")
        assert ready.status_code == 200 and ready.json()["current"], ready.text
        response = api.request("GET", f"/api/v1/statement-exports/{export['id']}/file")
        assert response.status_code == 200, response.text
        content = response.content
        sheet = load_workbook(BytesIO(content)).active
        assert "正式" in sheet["A1"].value
        assert sheet["G7"].value == 60 and sheet["I7"].value == 600
        assert approval["id"] in sheet["J7"].value
        items = api.request("GET", "/api/v1/archive-sources", params={"kind": "statement_export"})
        assert items.status_code == 200, items.text
        item = next(row for row in items.json()["items"] if row["source_id"] == export["id"])
        assert item["current"] and "正式" in item["title"]
        body = {
            "rule": "month_customer",
            "selections": [{"key": item["key"], "token": item["token"]}],
        }
        preview = api.request("POST", "/api/v1/archives/preview", json=body)
        assert preview.status_code == 200, preview.text
        archive = command(
            flow, "/api/v1/archives", body | {"preview_token": preview.json()["preview_token"]}
        )
        generated = api.request("POST", f"/api/v1/archives/{archive['id']}/generate")
        assert generated.status_code == 200 and generated.json()["status"] == "ready", (
            generated.text
        )
        archive_file = api.request("GET", f"/api/v1/archives/{archive['id']}/file")
        assert archive_file.status_code == 200, archive_file.text
        with ZipFile(BytesIO(archive_file.content)) as zipped:
            packages = [name for name in zipped.namelist() if name.endswith(".zip")]
            assert len(packages) == 1 and "/正式对账资料/" in packages[0]
            with ZipFile(BytesIO(zipped.read(packages[0]))) as materials:
                assert (
                    materials.read(
                        f"statements/{statement['id']}/v{formal['version']}/statement.xlsx"
                    )
                    == content
                )
        following = blank(flow, "2026-11-01", "2026-11-30")
        command(
            flow,
            f"/api/v1/statements/{following['id']}/claim-review",
            {
                "expected_version": 1,
                "review_id": review["id"],
                "review_version": review["version"],
                "quantity": "6",
                "unit_price": "10",
                "reason": "Synthetic partial resolution",
            },
        )
        assert detail(flow, statement["id"]) == formal
        assert (
            api.request("GET", f"/api/v1/statement-exports/{export['id']}/file").content == content
        )
        assert (
            api.request("GET", f"/api/v1/archives/{archive['id']}/file").content
            == archive_file.content
        )
    finally:
        application.state.services.close()
