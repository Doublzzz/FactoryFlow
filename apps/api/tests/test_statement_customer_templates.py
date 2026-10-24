from copy import deepcopy
from io import BytesIO

import pytest
from factoryflow.modules.statements.adapters.excel import render_excel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from test_statement_excel import snapshot


def template_file():
    book = Workbook()
    sheet = book.active
    sheet.title = "客户格式"
    sheet.append(["合成模板"])
    sheet.append(["样例客户", None, "样例期间", None, "样例状态"])
    sheet.append(["日期", "货号", "货物名称", "规格", "单位", "数量", "单价", "金额"])
    for row in range(4, 7):
        sheet.append(["旧日期", "旧货号", "旧名称", "旧规格", "旧单位", 999, 999, 999])
        sheet.cell(row, 3).fill = PatternFill("solid", fgColor="EAF0F6")
        sheet.row_dimensions[row].height = 25
    sheet.append(["已定价小计", None, None, None, None, None, None, "=SUM(H4:H6)"])
    sheet.append(["计税基数", None, None, None, None, None, None, 999])
    sheet.append(["税额", None, None, None, None, None, None, 999])
    sheet.append(["最终金额", None, None, None, None, None, None, 999])
    sheet.append(["调整说明", "旧调整"])
    sheet.append(["签收栏（合成）"])
    sheet.merge_cells("A1:H1")
    sheet.merge_cells("A12:H12")
    for col in "ABCDEFGH":
        sheet.column_dimensions[col].width = 15
    sheet.print_title_rows = "1:3"
    sheet.oddFooter.left.text = "合成固定页脚"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def mapping():
    return {
        "sheet_name": "客户格式",
        "detail_start_row": 4,
        "detail_end_row": 6,
        "columns": {
            "delivery_date": "A",
            "item_code": "B",
            "item_name": "C",
            "specification": "D",
            "unit": "E",
            "quantity": "F",
            "unit_price": "G",
            "amount": "H",
        },
        "cells": {
            "customer_name": "A2",
            "period": "C2",
            "status": "E2",
            "priced_subtotal": "H7",
            "tax_base": "H8",
            "tax_amount": "H9",
            "total": "H10",
            "adjustments": "B11",
        },
    }


def custom_snapshot(count=1):
    data = snapshot()
    data.update(
        template_id="customer-v1",
        template_version=1,
        customer_template={
            "id": "synthetic-template",
            "name": "合成客户格式",
            "version": 1,
            "mapping": mapping(),
        },
    )
    line = data["statement"]["lines"][0]
    data["statement"]["lines"] = [deepcopy(line) for _ in range(count)]
    data["statement"]["calculation"]["line_amounts"] = [None] * count
    return data


def test_customer_template_language_does_not_translate_cells_or_annotations():
    data = custom_snapshot(2)
    source = template_file()
    chinese = load_workbook(BytesIO(render_excel(data | {"locale": "zh-CN"}, source)))
    english = load_workbook(BytesIO(render_excel(data | {"locale": "en"}, source)))
    assert english.sheetnames == chinese.sheetnames == ["客户格式"]
    assert list(english.active.values) == list(chinese.active.values)
    assert english.active["A1"].value == "合成模板"
    assert "金额未确定" in english.active["E2"].value
    assert english.active.oddFooter.left.text == chinese.active.oddFooter.left.text
    assert str(english.active.print_area) == str(chinese.active.print_area)


@pytest.mark.parametrize("count", [0, 1, 3, 8, 301])
def test_customer_template_resizes_details_and_moves_all_footer_rows(count):
    sheet = load_workbook(BytesIO(render_excel(custom_snapshot(count), template_file()))).active
    assert sheet.cell(4 + count, 1).value == "已定价小计"
    assert sheet.cell(9 + count, 1).value == "签收栏（合成）"
    assert sheet.max_row == 9 + count
    assert f"A{9 + count}:H{9 + count}" in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet["A2"].data_type == "s"
    assert "金额未确定" in sheet["E2"].value
    assert sheet.cell(7 + count, 8).value is None
    assert not any(c.data_type == "f" for row in sheet for c in row)
    assert not any(str(c.value).startswith("旧") for row in sheet for c in row)
    assert sheet.print_title_rows == "$1:$3"
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.oddFooter.left.text == "合成固定页脚"
    if count:
        assert sheet["C4"].fill.fgColor.rgb == "00EAF0F6"
        assert sheet["G4"].value == "待定价"
        assert sheet["B4"].value == "=1+1"


def test_customer_template_preserves_precise_values_and_expands_long_text():
    data = custom_snapshot()
    data["statement"]["calculation"]["total"] = "12345678901234567890.12"
    data["statement"]["lines"][0]["source"]["specification"] = "第一行\n第二行\n第三行\n第四行"
    sheet = load_workbook(BytesIO(render_excel(data, template_file()))).active
    assert sheet["H8"].value == "12345678901234567890.12"
    assert sheet.row_dimensions[4].height >= 64


def altered_file(change):
    book = load_workbook(BytesIO(template_file()))
    change(book.active)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    "change,code",
    [
        (lambda s: s.__setitem__("A12", "=H10"), "STATEMENT_TEMPLATE_FORMULA"),
        (lambda s: s.merge_cells("C4:C5"), "STATEMENT_TEMPLATE_MAPPING"),
        (lambda s: s.merge_cells("F4:G4"), "STATEMENT_TEMPLATE_MAPPING"),
        (lambda s: setattr(s.row_dimensions[2], "hidden", True), "STATEMENT_TEMPLATE_MAPPING"),
        (lambda s: setattr(s.column_dimensions["F"], "hidden", True), "STATEMENT_TEMPLATE_MAPPING"),
        (lambda s: setattr(s.protection, "sheet", True), "STATEMENT_TEMPLATE_UNSUPPORTED"),
        (lambda s: s.__setitem__("AE1", "outside"), "STATEMENT_TEMPLATE_UNSUPPORTED"),
        (lambda s: s.__setitem__("A301", "outside"), "STATEMENT_TEMPLATE_UNSUPPORTED"),
        (
            lambda s: (
                setattr(s.auto_filter, "ref", "A3:H6"),
                s.auto_filter.add_filter_column(0, ["合成筛选条件"]),
            ),
            "STATEMENT_TEMPLATE_UNSUPPORTED",
        ),
        (
            lambda s: setattr(s["A12"], "hyperlink", "https://example.invalid/"),
            "STATEMENT_TEMPLATE_UNSUPPORTED",
        ),
    ],
)
def test_template_rejects_unpreserved_features_and_unsafe_bindings(change, code):
    from factoryflow.kernel.contracts import BusinessError
    from factoryflow.modules.statements.adapters.customer_excel import validate_mapping

    with pytest.raises(BusinessError) as caught:
        validate_mapping(altered_file(change), mapping())
    assert caught.value.code == code


@pytest.mark.parametrize(
    "change",
    [
        lambda m: m["columns"].update(quantity="G"),
        lambda m: m["cells"].update(status="A2"),
        lambda m: m["cells"].update(status="E4"),
        lambda m: m["cells"].update(total="H12"),
        lambda m: m["columns"].pop("unit"),
    ],
)
def test_template_rejects_ambiguous_or_incomplete_mapping(change):
    from factoryflow.kernel.contracts import BusinessError
    from factoryflow.modules.statements.adapters.customer_excel import validate_mapping

    chosen = mapping()
    change(chosen)
    # Without subtotal, total must be adjacent to details.
    if chosen["cells"].get("total") == "H12":
        chosen["cells"].pop("priced_subtotal")
    with pytest.raises(BusinessError):
        validate_mapping(template_file(), chosen)


def test_template_formula_outputs_and_formats_are_replaced_by_saved_values():
    content = altered_file(lambda s: setattr(s["F4"], "number_format", "0%"))
    data = custom_snapshot()
    sheet = load_workbook(BytesIO(render_excel(data, content))).active
    assert sheet["F4"].value == 12.125
    assert sheet["F4"].number_format == "0.####"


def test_customer_template_keeps_range_column_widths_and_visible_dark_fill_bindings():
    from openpyxl.styles import Font

    def change(sheet):
        sheet.column_dimensions.clear()
        sheet.column_dimensions["A"].min = 1
        sheet.column_dimensions["A"].max = 8
        sheet.column_dimensions["A"].width = 20
        sheet["E2"].fill = PatternFill("solid", fgColor="172B4D")
        sheet["E2"].font = Font(color="FFFFFF")
        sheet["H10"].fill = PatternFill("solid", fgColor="172B4D")

    sheet = load_workbook(BytesIO(render_excel(custom_snapshot(), altered_file(change)))).active
    assert len(sheet.column_dimensions) == 1
    assert sheet.column_dimensions["A"].min == 1
    assert sheet.column_dimensions["A"].max == 8
    assert sheet.column_dimensions["A"].width == 20
    assert sheet["E2"].font.color.rgb[-6:] == "FFFFFF"
    assert sheet["H8"].font.color.rgb[-6:] == "FFFFFF"


def test_business_bindings_replace_fill_alignment_and_indent():
    from openpyxl.styles import Alignment

    def change(sheet):
        for address in ("H4", "H10", "E2"):
            sheet[address].alignment = Alignment(horizontal="fill", indent=10, relativeIndent=8)

    data = custom_snapshot()
    data["statement"]["lines"][0].update(quantity="12", unit_price="1")
    data["statement"]["calculation"].update(
        line_amounts=["12.00"],
        priced_subtotal="12.00",
        tax_base="12.00",
        tax_amount="0.00",
        total="12.00",
        warnings=[],
    )
    sheet = load_workbook(BytesIO(render_excel(data, altered_file(change)))).active
    for address, horizontal in (("H4", "right"), ("H8", "right"), ("E2", "left")):
        alignment = sheet[address].alignment
        assert alignment.horizontal == horizontal
        assert alignment.indent == 0 and alignment.relativeIndent == 0


@pytest.mark.parametrize("theme,expected", [(0, "000000"), (1, "FFFFFF")])
def test_theme_fill_has_readable_business_text(theme, expected):
    from openpyxl.styles import Color

    def change(sheet):
        sheet["H10"].fill = PatternFill("solid", fgColor=Color(theme=theme))

    sheet = load_workbook(BytesIO(render_excel(custom_snapshot(), altered_file(change)))).active
    assert sheet["H8"].font.color.rgb[-6:] == expected


@pytest.mark.parametrize("field", ["specification", "adjustments"])
def test_unprintable_long_bindings_fail_instead_of_clipping(field):
    from factoryflow.kernel.contracts import BusinessError

    data = custom_snapshot()
    if field == "specification":
        data["statement"]["lines"][0]["source"][field] = "合成一行\n" * 30
    else:
        data["statement"]["adjustments"] = [
            {"stage": "after_tax", "kind": "合成调整", "amount": "1", "reason": "合成一行\n" * 30}
        ]
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_TEXT_LIMIT"):
        render_excel(data, template_file())


def test_template_tax_adjustments_and_formal_settlement_rows_are_explicit():
    data = custom_snapshot()
    statement = data["statement"]
    statement["status"] = "confirmed"
    statement["tax_mode"] = "add_tax"
    statement["adjustments"] = [
        {"stage": "before_tax", "kind": "合成扣减", "amount": "-2", "reason": "合成说明"}
    ]
    statement["settlement_lines"], statement["lines"] = statement["lines"], []
    statement["settlement_lines"][0]["unit_price"] = "1.2"
    statement["calculation"].update(
        line_amounts=["14.55"],
        priced_subtotal="14.55",
        tax_base="12.55",
        tax_amount="1.63",
        total="14.18",
        warnings=[],
    )
    sheet = load_workbook(BytesIO(render_excel(data, template_file()))).active
    assert sheet["E2"].value.startswith("正式 v1")
    assert "未税加税" in sheet["E2"].value and "CNY" in sheet["E2"].value
    assert sheet["H4"].value == 14.55
    assert sheet["H5"].value == 14.55
    assert sheet["H6"].value == 12.55
    assert sheet["H7"].value == 1.63
    assert sheet["H8"].value == 14.18
    assert "税前 / 合成扣减 / -2：合成说明" in sheet["B9"].value


def test_template_rejects_unexpressed_adjustments_and_excess_output():
    from factoryflow.kernel.contracts import BusinessError

    data = custom_snapshot()
    data["statement"]["adjustments"] = [
        {"stage": "after_tax", "kind": "合成运费", "amount": "2", "reason": "合成说明"}
    ]
    data["customer_template"]["mapping"]["cells"].pop("adjustments")
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_COVERAGE"):
        render_excel(data, template_file())
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_OUTPUT_LIMIT"):
        render_excel(custom_snapshot(1001), template_file())


@pytest.mark.parametrize(
    "replacement",
    [
        b'<!DOCTYPE worksheet [<!ENTITY x "bad">]><worksheet />',
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:XFD1048576" /></worksheet>',
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<drawing /></worksheet>",
    ],
)
def test_template_rejects_dangerous_xml_before_loading(replacement):
    from zipfile import ZipFile

    from factoryflow.kernel.contracts import BusinessError
    from factoryflow.modules.statements.adapters.customer_excel import inspect_template

    output = BytesIO()
    with ZipFile(BytesIO(template_file())) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            target.writestr(
                name, replacement if name == "xl/worksheets/sheet1.xml" else source.read(name)
            )
    with pytest.raises(BusinessError):
        inspect_template(output.getvalue())
