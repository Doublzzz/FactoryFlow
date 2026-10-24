"""Synthetic packaging/layout regressions; never store a customer workbook here."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from factoryflow.kernel.contracts import BusinessError
from factoryflow.modules.statements.adapters.customer_excel import inspect_template
from factoryflow.modules.statements.adapters.excel import render_excel
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.packaging.custom import StringProperty
from test_statement_customer_templates import custom_snapshot, mapping, template_file

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CALC = "{http://schemas.microsoft.com/office/spreadsheetml/2018/calcfeatures}"
WPS = "{http://www.wps.cn/officeDocument/2017/etCustomData}"


def edited_book(change):
    book = load_workbook(BytesIO(template_file()))
    change(book)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def edited_xml(content, filename, change, directories=()):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for directory in directories:
            target.writestr(directory, b"")
        for name in source.namelist():
            data = source.read(name)
            if name == filename:
                root = ET.fromstring(data)
                change(root)
                data = ET.tostring(root)
            target.writestr(name, data)
    return output.getvalue()


def calculation_flags(root):
    extension = ET.SubElement(ET.SubElement(root, NS + "extLst"), NS + "ext")
    extension.set("uri", "{B58B0392-4F1F-4190-BB64-5DF3571DCE5F}")
    ET.SubElement(
        ET.SubElement(extension, CALC + "calcFeatures"), CALC + "feature", name="microsoft.com:RD"
    )


@pytest.mark.parametrize(
    "feature", ["directories", "properties", "calculation_flags", "empty_sheet"]
)
def test_common_office_packaging_preserves_one_content_sheet(feature):
    content = template_file()
    if feature == "directories":
        content = edited_xml(
            content, "xl/workbook.xml", lambda _: None, ("xl/", "xl/worksheets/", "docProps/")
        )
    elif feature == "properties":
        content = edited_book(
            lambda b: b.custom_doc_props.append(
                StringProperty(name="Synthetic", value="Synthetic metadata")
            )
        )
    elif feature == "calculation_flags":
        content = edited_xml(content, "xl/workbook.xml", calculation_flags)
    else:

        def blank(book):
            sheet = book.create_sheet("空白备用页")
            sheet.column_dimensions["A"].width = 22
            sheet.cell(50, 4).number_format = "0.00"
            book.active = 1

        content = edited_book(blank)
    assert inspect_template(content)["sheet_name"] == mapping()["sheet_name"]
    output = load_workbook(BytesIO(render_excel(custom_snapshot(), content)))
    if feature == "properties":
        assert output.custom_doc_props["Synthetic"].value == "Synthetic metadata"
    if feature == "empty_sheet":
        assert len(output.worksheets) == 2
        assert output.active.title == mapping()["sheet_name"]
        assert output.worksheets[1].column_dimensions["A"].width == 22
        assert output.worksheets[1]["D50"].number_format == "0.00"


@pytest.mark.parametrize("count", [0, 1, 12])
def test_unused_filter_moves_to_actual_detail_region(count):
    content = edited_book(lambda b: setattr(b.active.auto_filter, "ref", "A1:H60"))

    def empty_extension(root):
        node = root.find(NS + "autoFilter")
        node.set(WPS + "filterBottomFollowUsedRange", "0")
        ET.SubElement(node, NS + "extLst")

    content = edited_xml(content, "xl/worksheets/sheet1.xml", empty_extension)
    sheet = load_workbook(BytesIO(render_excel(custom_snapshot(count), content))).active
    assert sheet.auto_filter.ref == (f"A3:H{3 + count}" if count else None)


@pytest.mark.parametrize(
    "kind", ["unknown_extension", "filter_criteria", "extra_content", "unsafe_directory"]
)
def test_compatibility_does_not_accept_unknown_or_lossy_features(kind):
    content = template_file()
    if kind == "unknown_extension":

        def extension(root):
            calculation_flags(root)
            root.find(NS + "extLst")[0].set("uri", "unknown")

        content = edited_xml(content, "xl/workbook.xml", extension)
    elif kind == "filter_criteria":

        def filtered(book):
            book.active.auto_filter.ref = "A3:H6"
            book.active.auto_filter.add_filter_column(0, ["synthetic"])

        content = edited_book(filtered)
    elif kind == "extra_content":
        content = edited_book(lambda b: b.create_sheet("另一份内容").append(["合成内容"]))
    else:
        content = edited_xml(content, "xl/workbook.xml", lambda _: None, ("../",))
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_UNSUPPORTED"):
        inspect_template(content)


def compact_case():
    def compact(book):
        sheet = book.active
        sheet["C3"] = "品名 / 规格"
        sheet.merge_cells("C2:H2")

    content = edited_book(compact)
    data = custom_snapshot()
    fields = data["customer_template"]["mapping"]
    fields["columns"].pop("item_name")
    fields["columns"].pop("specification")
    fields["columns"]["item_name_specification"] = "C"
    fields["cells"].pop("status")
    fields["cells"].pop("tax_base")
    fields["cells"].pop("tax_amount")
    return content, data


def test_combined_item_and_period_status_keep_all_business_information():
    content, data = compact_case()
    fields = inspect_template(content)["suggested_mapping"]["columns"]
    assert fields["item_name_specification"] == "C"
    sheet = load_workbook(BytesIO(render_excel(data, content))).active
    assert sheet["C4"].value == "测试零件 / synthetic specification"
    assert "2026-10-01 至 2026-10-31" in sheet["C2"].value
    assert "草稿 v1" in sheet["C2"].value and "金额未确定" in sheet["C2"].value


@pytest.mark.parametrize("missing_price", [False, True])
def test_compact_tax_template_discloses_saved_base_and_tax_in_header(missing_price):
    content, data = compact_case()
    statement = data["statement"]
    statement.update(tax_mode="add_tax", tax_rate="9")
    if not missing_price:
        statement["lines"][0].update(quantity="2", unit_price="6")
        statement["calculation"].update(
            line_amounts=["12.00"],
            priced_subtotal="12.00",
            tax_base="12.00",
            tax_amount="1.08",
            total="13.08",
            warnings=[],
        )
    sheet = load_workbook(BytesIO(render_excel(data, content))).active
    assert (
        "计税基数：待确定" in sheet["C2"].value
        if missing_price
        else "计税基数：12.00" in sheet["C2"].value
    )
    assert (
        "税额：待确定" in sheet["C2"].value if missing_price else "税额：1.08" in sheet["C2"].value
    )
    assert sheet["H8"].value == (None if missing_price else 13.08)


def test_mapped_order_and_delivery_references_use_saved_sources():
    data = custom_snapshot()
    data["statement"]["lines"][0]["source"].update(
        order_number="SYNTHETIC-ORDER", delivery_id="synthetic-delivery"
    )
    fields = data["customer_template"]["mapping"]["columns"]
    fields.pop("item_code")
    fields.pop("specification")
    fields.update(order_number="B", delivery_reference="D")
    sheet = load_workbook(BytesIO(render_excel(data, template_file()))).active
    assert sheet["B4"].value == "SYNTHETIC-ORDER"
    assert sheet["D4"].value == "synthetic-delivery"
    data["statement"]["lines"][0]["source"].pop("order_number")
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_COVERAGE"):
        render_excel(data, template_file())


def test_rich_static_header_format_survives_without_running_formulas():
    content = edited_book(
        lambda b: setattr(
            b.active["A1"],
            "value",
            CellRichText("合成", TextBlock(InlineFont(b=True, color="AA0000"), "富文本标题")),
        )
    )
    output = load_workbook(BytesIO(render_excel(custom_snapshot(), content)), rich_text=True)
    assert isinstance(output.active["A1"].value, CellRichText)
    assert output.active["A1"].value[1].font.b is True
    assert output.active["A1"].value[1].font.color.rgb[-6:] == "AA0000"


def test_unused_builtin_slicer_style_metadata_is_accepted_but_other_styles_are_not():
    def extension(root):
        node = ET.SubElement(
            ET.SubElement(root, NS + "extLst"),
            NS + "ext",
            uri="{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}",
        )
        ET.SubElement(
            node,
            "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}slicerStyles",
            defaultSlicerStyle="SlicerStyleLight1",
        )

    content = edited_xml(template_file(), "xl/styles.xml", extension)
    assert inspect_template(content)["rows"] == 12
    content = edited_xml(
        content, "xl/styles.xml", lambda r: r.find(NS + "extLst")[0].set("uri", "unknown")
    )
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_UNSUPPORTED"):
        inspect_template(content)


@pytest.mark.parametrize(
    "mode,label", [("add_tax", "含税"), ("included", "含税"), ("untaxed_final", "未税")]
)
def test_total_tax_label_follows_statement_instead_of_static_template(mode, label):
    content = edited_book(lambda b: setattr(b.active["G10"], "value", "旧税制标签"))
    data = custom_snapshot()
    data["statement"]["tax_mode"] = mode
    data["customer_template"]["mapping"]["cells"]["total_label"] = "G10"
    sheet = load_workbook(BytesIO(render_excel(data, content))).active
    assert sheet["G8"].value == label


def test_print_width_ignores_outside_empty_formatting_and_repeats_detail_header():
    def layout(book):
        sheet = book.active
        sheet.cell(12, 9).number_format = "0.00"
        sheet.print_area = "A1:H90"
        sheet.print_title_rows = "1:2"

    content = edited_book(layout)
    sheet = load_workbook(BytesIO(render_excel(custom_snapshot(12), content))).active
    assert sheet.print_area.endswith("$A$1:$H$21")
    assert sheet.print_title_rows == "$1:$3"


def test_unmodified_header_height_is_preserved_and_whole_numbers_have_no_decimal_dot():
    from openpyxl.styles import Alignment

    def layout(book):
        book.active["B3"] = "合成原表头\n合成原表头"
        book.active["B3"].alignment = Alignment(wrap_text=True)
        book.active.row_dimensions[3].height = 24

    data = custom_snapshot()
    data["statement"]["lines"][0].update(quantity="2", unit_price="6")
    sheet = load_workbook(BytesIO(render_excel(data, edited_book(layout)))).active
    assert sheet.row_dimensions[3].height == 24
    assert sheet["F4"].number_format == "0"
    assert sheet["G4"].number_format == "0"


def test_bound_merged_header_cannot_extend_outside_print_area():
    content, data = compact_case()
    book = load_workbook(BytesIO(content))
    book.active.unmerge_cells("C2:H2")
    book.active.merge_cells("C2:J2")
    book.active.print_area = "A1:H90"
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_MAPPING"):
        render_excel(data, stream.getvalue())


@pytest.mark.parametrize("value", [0, False])
def test_zero_and_false_are_content_in_a_second_sheet(value):
    content = edited_book(lambda b: b.create_sheet("合成第二页").append([value]))
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_UNSUPPORTED"):
        inspect_template(content)


@pytest.mark.parametrize("status", ["draft", "confirmed"])
def test_compact_period_retains_history_status_and_missing_price_warning(status):
    content, data = compact_case()
    data["statement"]["status"] = status
    data["current"] = False
    sheet = load_workbook(BytesIO(render_excel(data, content))).active
    assert "历史版本" in sheet["C2"].value
    assert ("正式 v1" if status == "confirmed" else "草稿 v1") in sheet["C2"].value
    assert "金额未确定（缺单价）" in sheet["C2"].value
