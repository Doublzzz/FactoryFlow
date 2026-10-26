"""Owned 0014/current backups: exact old evidence, retry identity and USD restoration."""

from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import psycopg
import pytest
from alembic import command as alembic_command
from alembic.config import Config
from conftest import API, ROOT, SYNTHETIC_PASSWORD
from factoryflow.bootstrap import build_services, create_app
from factoryflow.modules.statements.adapters.repository import StatementRepository
from factoryflow.platform import database
from factoryflow.platform.backup import restore_backup, validate_backup
from factoryflow.platform.config import Settings
from factoryflow.platform.database import engine_for
from factoryflow.platform.migrations import upgrade_database
from openpyxl import load_workbook
from psycopg import sql
from sqlalchemy import text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError
from test_archive_migration import BINARY, owned_database
from test_postgres_currencies import delivered, draft_input
from test_postgres_statement_lifecycle import approve, command, reserve, settle, setup_statement


def frozen_rows(engine):
    names = ("confirmed_revisions", "statement_events", "statement_exports", "command_results")
    with engine.connect() as conn:
        return {
            name: sorted(conn.execute(text(f"SELECT to_jsonb(t)::text FROM {name} t")).scalars())
            for name in names
        }


def test_0014_backup_rollback_upgrade_keep_old_cny_snapshots_and_replay(pg, tmp_path, monkeypatch):
    with owned_database(pg, "ff_currency_old_") as (name, admin):
        with psycopg.connect(admin.render_as_string(hide_password=False), autocommit=True) as conn:
            conn.execute(sql.SQL("CREATE DATABASE {} OWNER ff_admin").format(sql.Identifier(name)))
        owner_url = pg.admin_url.set(database=name).render_as_string(hide_password=False)
        owner = engine_for(owner_url)
        app = None
        try:
            config = Config(str(ROOT / "alembic.ini"))
            with owner.begin() as conn:
                config.attributes["connection"] = conn
                alembic_command.upgrade(config, "0014_finished_goods_inventory")
            objects = tmp_path / "old-objects"
            objects.mkdir()
            app = create_app(settings=Settings(owner_url, objects))
            runtime = app.state.services
            actor = runtime.identity.provision_factory(
                "Synthetic Legacy USD", "synthetic", SYNTHETIC_PASSWORD
            )
            customer = runtime.customers.create(actor, uuid4(), "Synthetic Legacy Customer")
            api = API(app)
            fixture = SimpleNamespace(api=api, customer_id=customer["id"])
            original_get = StatementRepository.get

            def old_projection(self, *args, **kwargs):
                value = original_get(self, *args, **kwargs)
                value.pop("currency", None)
                return value

            # Seed exactly the old no-header-currency shape, without disabling evidence guards.
            with monkeypatch.context() as patch:
                patch.setattr(database, "REVISION", "0014_finished_goods_inventory")
                patch.setattr(StatementRepository, "get", old_projection)
                api.login("synthetic")
                statement, source, sent, feedback = setup_statement(fixture)
                approval = approve(fixture, statement, source, sent, feedback, "100")
                reserve(fixture, statement, approval, "100")
                formal = settle(fixture, statement)
                export = runtime.statement_exports.create(
                    actor, uuid4(), UUID(statement["id"]), formal["version"]
                )
                assert runtime.statement_exports.generate(actor, export["id"])["status"] == "ready"
                key = uuid4()
                body = draft_input(fixture, "CNY", start="2026-11-01", end="2026-11-30")
                old_result = runtime.statements.create(actor, key, body)
                assert "currency" not in old_result
                pending = runtime.statement_exports.create(
                    actor, uuid4(), UUID(statement["id"]), formal["version"]
                )
            before = frozen_rows(owner)
            with owner.connect() as conn:
                assert not conn.scalar(
                    text("SELECT bool_or(snapshot ? 'currency') FROM statement_events")
                )
            excel = runtime.statement_exports.download(actor, export["id"])[0]
            original_objects = {path.name: path.read_bytes() for path in objects.glob("*.json")}
            backup = tmp_path / "old-backup"
            manifest = runtime.backup(backup, BINARY)
            assert manifest["schema_revision"] == "0014_finished_goods_inventory"
            assert validate_backup(backup) == manifest
            with owner.begin() as conn:
                conn.execute(
                    text(
                        "CREATE FUNCTION guard_statement_currency() RETURNS trigger "
                        "LANGUAGE plpgsql AS $$ BEGIN RETURN NEW; END; $$"
                    )
                )
            with pytest.raises(DBAPIError):
                upgrade_database(owner, ROOT)
            with owner.begin() as conn:
                assert (
                    conn.scalar(text("SELECT version_num FROM alembic_version"))
                    == "0014_finished_goods_inventory"
                )
                assert not conn.scalar(
                    text(
                        "SELECT EXISTS(SELECT 1 FROM information_schema.columns "
                        "WHERE table_name='statements' AND column_name='currency')"
                    )
                )
                assert conn.scalar(text("SELECT to_regprocedure('guard_order_currency()')")) is None
                conn.execute(text("DROP FUNCTION guard_statement_currency()"))
            assert frozen_rows(owner) == before
            upgrade_database(owner, ROOT)
            runtime.check_ready()
            assert frozen_rows(owner) == before
            assert runtime.statements.create(actor, key, body) == old_result
            assert runtime.statements.detail(actor, UUID(statement["id"]))["currency"] == "CNY"
            assert runtime.statement_exports.download(actor, export["id"])[0] == excel
            with owned_database(pg, "ff_currency_restore_") as (restored_name, restore_admin):
                restored_objects = tmp_path / "restored-objects"
                restore_backup(restore_admin, restored_name, restored_objects, backup, BINARY, ROOT)
                url = make_url(pg.settings.database_url).set(database=restored_name)
                restored = build_services(
                    Settings(url.render_as_string(hide_password=False), restored_objects)
                )
                try:
                    restored.check_ready()
                    assert frozen_rows(restored.engine) == before
                    assert restored.statements.create(actor, key, body) == old_result
                    assert (
                        restored.statements.detail(actor, UUID(statement["id"]))["currency"]
                        == "CNY"
                    )
                    assert restored.statement_exports.download(actor, export["id"])[0] == excel
                    for filename, content in original_objects.items():
                        assert (restored_objects / filename).read_bytes() == content
                    assert (
                        restored.statement_exports.generate(actor, pending["id"])["status"]
                        == "ready"
                    )
                    sheet = load_workbook(
                        BytesIO(restored.statement_exports.download(actor, pending["id"])[0])
                    ).active
                    assert "CNY" in sheet["A3"].value
                finally:
                    restored.close()
        finally:
            if app is not None:
                app.state.services.close()
            owner.dispose()


def test_usd_backup_restores_english_export_order_balance_and_retry(pg, tmp_path):
    order, delivery = delivered(pg, "USD")
    body, key = draft_input(pg, "USD", delivery), uuid4()
    statement = command(pg, "/api/v1/statements", body, key)
    export = pg.runtime.statement_exports.create(
        pg.actor, uuid4(), UUID(statement["id"]), 1, locale="en"
    )
    assert pg.runtime.statement_exports.generate(pg.actor, export["id"])["status"] == "ready"
    original = pg.runtime.statement_exports.download(pg.actor, export["id"])[0]
    backup = tmp_path / "usd-backup"
    manifest = pg.runtime.backup(backup, BINARY)
    assert validate_backup(backup) == manifest
    with owned_database(pg, "ff_currency_usd_") as (name, admin):
        objects = tmp_path / "usd-restored"
        restore_backup(admin, name, objects, backup, BINARY, ROOT)
        url = make_url(pg.settings.database_url).set(database=name)
        runtime = build_services(Settings(url.render_as_string(hide_password=False), objects))
        try:
            runtime.check_ready()
            assert runtime.statements.create(pg.actor, key, body) == statement
            current = runtime.orders.detail(pg.actor, UUID(order["id"]))
            assert (
                current["currency"] == "USD" and current["lines"][0]["remaining_quantity"] == "70"
            )
            assert runtime.statement_exports.download(pg.actor, export["id"])[0] == original
            sheet = load_workbook(BytesIO(original)).active
            assert "USD" in sheet["A3"].value
            assert "Total (USD)" in [cell.value for row in sheet for cell in row]
        finally:
            runtime.close()
