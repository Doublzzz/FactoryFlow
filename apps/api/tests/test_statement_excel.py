from io import BytesIO

import pytest
from factoryflow.modules.statements.adapters.excel import render_excel
from openpyxl import load_workbook


def snapshot():
    return {
        "id": "synthetic-export",
        "template_id": "standard-v1",
        "template_version": 1,
        "statement": {
            "id": "synthetic-statement",
            "version": 1,
            "customer_name": "=SYNTHETIC()",
            "period_start": "2026-10-01",
            "period_end": "2026-10-31",
            "tax_mode": "untaxed_final",
            "tax_rate": "0",
            "adjustments": [],
            "lines": [
                {
                    "status": "included",
                    "quantity": "12.125",
                    "unit_price": None,
                    "reason": "synthetic",
                    "delivery_id": "synthetic-delivery",
                    "source": {
                        "item_code": "=1+1",
                        "item_name": "测试零件",
                        "specification": "synthetic specification",
                        "description": "",
                        "delivery_date": "2026-10-01",
                        "canonical_unit": "KG",
                    },
                }
            ],
            "calculation": {
                "line_amounts": [None],
                "priced_subtotal": "0.00",
                "tax_base": None,
                "tax_amount": None,
                "total": None,
                "warnings": ["PRICE_REQUIRED"],
            },
        },
    }


def test_excel_missing_price_is_not_zero_and_text_is_not_formula():
    book = load_workbook(BytesIO(render_excel(snapshot())), data_only=False)
    cells = [cell for row in book.active for cell in row]
    assert not any(cell.data_type == "f" for cell in cells)
    assert any(cell.value == "=1+1" and cell.data_type == "s" for cell in cells)
    assert any(cell.value == "待定价" for cell in cells)
    total = next(cell for cell in cells if cell.value == "最终金额（CNY）")
    assert book.active.cell(total.row, 10).value is None


def test_excel_large_money_keeps_exact_digits():
    data = snapshot()
    data["statement"]["calculation"]["total"] = "12345678901234567890.12"
    book = load_workbook(BytesIO(render_excel(data)))
    assert any(c.value == "12345678901234567890.12" for row in book.active for c in row)


@pytest.mark.parametrize("template_id", ["standard-v1", "delivery-v1"])
def test_english_snapshot_preserves_data_missing_price_and_decimal_values(template_id):
    data = snapshot() | {"locale": "en", "template_id": template_id}
    data["statement"]["lines"][0]["external_document_number"] = "原号-EN-SYN"
    sheet = load_workbook(BytesIO(render_excel(data))).active
    values = [c.value for row in sheet for c in row]
    assert sheet.title == "Statement Draft"
    assert "Price pending" in values
    assert "测试零件" in values
    assert "=1+1" in values
    assert any("External number: 原号-EN-SYN" in str(v) for v in values)
    assert not any(c.data_type == "f" for row in sheet for c in row)
    total = next(c for row in sheet for c in row if c.value == "Total (CNY)")
    assert sheet.cell(total.row, 10).value is None
    assert sheet.row_dimensions[6].height >= 36
    assert sheet.page_setup.fitToWidth == 1


def test_language_is_explicit_and_old_snapshots_remain_chinese():
    legacy = load_workbook(BytesIO(render_excel(snapshot()))).active
    explicit = load_workbook(BytesIO(render_excel(snapshot() | {"locale": "zh-CN"}))).active
    assert list(legacy.values) == list(explicit.values)
    assert legacy.title == "对账草稿"
    with pytest.raises(ValueError, match="locale"):
        render_excel(snapshot() | {"locale": "fr"})


def test_english_historical_final_label_fits_excel_sheet_title():
    data = snapshot() | {"locale": "en", "current": False}
    data["statement"]["status"] = "confirmed"
    sheet = load_workbook(BytesIO(render_excel(data))).active
    assert sheet.title == "Statement Historical Final"
    assert len(sheet.title) <= 31
    assert "Historical" in sheet["A1"].value


def test_explicit_specification_line_breaks_receive_printable_height():
    data = snapshot()
    data["statement"]["lines"][0]["source"]["specification"] = "第一行\n第二行\n第三行\n第四行"
    book = load_workbook(BytesIO(render_excel(data)))
    assert book.active.row_dimensions[7].height >= 64
    assert book.active["E7"].value == "第一行\n第二行\n第三行\n第四行"


@pytest.mark.parametrize("template_id", ["standard-v1", "delivery-v1"])
def test_builtin_exports_use_portrait_pages_and_keep_long_tables_paginated(template_id):
    data = snapshot()
    data["template_id"] = template_id
    line = data["statement"]["lines"][0]
    data["statement"]["lines"] = [dict(line, reason=f"synthetic row {i}") for i in range(50)]
    data["statement"]["calculation"]["line_amounts"] = [None] * 50

    sheet = load_workbook(BytesIO(render_excel(data))).active

    assert sheet.page_setup.orientation == "portrait"
    assert sheet.page_setup.paperSize == 9
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.print_title_rows == "$1:$6"
    details = [row for row in sheet if str(row[9].value).startswith("synthetic row ")]
    assert len(details) == 50
    assert details[-1][9].value == "synthetic row 49"
    subtotal = next(cell for row in sheet for cell in row if cell.value == "已定价小计")
    assert subtotal.row == details[-1][0].row + 1
