"""Synthetic customer references remain metadata and fixed lifecycle evidence."""

from io import BytesIO
from uuid import uuid4

import pytest
from factoryflow.kernel.contracts import BusinessError
from factoryflow.modules.statements.adapters.excel import render_excel
from factoryflow.modules.statements.domain.customer_templates import require_coverage
from factoryflow.modules.statements.domain.draft import prepare_draft
from openpyxl import load_workbook
from test_statement_excel import snapshot


def draft(number="SYN-OUT-001"):
    return {
        "customer_id": str(uuid4()),
        "period_start": "2026-10-01",
        "period_end": "2026-10-31",
        "tax_mode": "untaxed_final",
        "tax_rate": "0",
        "lines": [
            {
                "delivery_id": str(uuid4()),
                "order_line_id": str(uuid4()),
                "quantity": "1",
                "external_document_number": number,
            }
        ],
    }


@pytest.mark.parametrize("number", ["", "  SYN-RETURN-001  ", "=1+1", "x" * 120])
def test_reference_normalization(number):
    assert (
        prepare_draft(draft(number))["lines"][0].get("external_document_number", "")
        == number.strip()
    )


def test_empty_reference_keeps_legacy_command_fingerprint():
    payload = draft("")
    legacy = {
        **payload,
        "lines": [
            {k: v for k, v in payload["lines"][0].items() if k != "external_document_number"}
        ],
    }
    canonical = prepare_draft(legacy)
    assert "external_document_number" not in canonical["lines"][0]
    assert prepare_draft(payload) == canonical


@pytest.mark.parametrize("number", [None, 123, "x" * 121, "SYN\n1", "SYN\x7f1", "SYN\x001"])
def test_reference_invalid(number):
    with pytest.raises(BusinessError, match="STATEMENT_EXTERNAL_NUMBER_INVALID"):
        prepare_draft(draft(number))


def test_mapped_reference_required_without_uuid_fallback():
    data = snapshot()["statement"]
    mapping = {"cells": {}, "columns": {"external_document_number": "A"}}
    with pytest.raises(BusinessError, match="STATEMENT_TEMPLATE_COVERAGE"):
        require_coverage(mapping, data)
    data["lines"][0]["external_document_number"] = "SYN-OUT-001"
    require_coverage(mapping, data)


def test_standard_excel_preserves_reference_as_safe_text():
    data = snapshot()
    data["statement"]["lines"][0]["external_document_number"] = "=SYNTHETIC()"
    book = load_workbook(BytesIO(render_excel(data)), data_only=False)
    cells = [cell for row in book.active for cell in row]
    assert any(
        "=SYNTHETIC()" in str(cell.value) and "外部单号" in str(cell.value) for cell in cells
    )
    assert not any(cell.data_type == "f" for cell in cells)


def save_number(pg, statement, number, *, key=None, expected=None):
    from test_postgres_statement_lifecycle import detail

    current = detail(pg, statement["id"])
    return pg.api.command(
        f"/api/v1/statements/{statement['id']}/save",
        {
            k: current[k]
            for k in (
                "customer_id",
                "period_start",
                "period_end",
                "tax_mode",
                "tax_rate",
                "adjustments",
            )
        }
        | {
            "expected_version": current["version"] if expected is None else expected,
            "reason": "Synthetic number entry",
            "lines": [
                {
                    k: line[k]
                    for k in ("delivery_id", "order_line_id", "quantity", "unit_price", "reason")
                }
                | {"external_document_number": number}
                for line in current["lines"]
                if line["status"] == "included"
            ],
        },
        key,
    )


def send_current(pg, statement):
    from test_postgres_statement_lifecycle import command, detail

    version = detail(pg, statement["id"])["version"]
    export = command(
        pg, f"/api/v1/statements/{statement['id']}/exports", {"expected_version": version}
    )
    response = pg.api.request("POST", f"/api/v1/statement-exports/{export['id']}/generate")
    assert response.json()["status"] == "ready", response.text
    sent = command(
        pg,
        f"/api/v1/statements/{statement['id']}/sent",
        {
            "expected_version": version,
            "export_id": export["id"],
            "person": "Synthetic sender",
            "occurred_on": "2026-10-04",
            "note": "Synthetic sent reference",
        },
    )
    return export, sent


def test_pg_number_evidence_review_and_formal_propagation(pg):
    from sqlalchemy import text
    from sqlalchemy.exc import DBAPIError
    from test_postgres_statement_lifecycle import (
        approve,
        blank,
        command,
        detail,
        reserve,
        settle,
        setup_statement,
    )

    statement, source, old_sent, feedback = setup_statement(pg)
    saved = save_number(pg, statement, "SYN-EXT-001")
    assert saved.status_code == 200, saved.text
    saved = saved.json()
    assert saved["lines"][0]["source"] == statement["lines"][0]["source"]
    assert saved["history"][0]["snapshot"]["lines"][0]["external_document_number"] == ""
    mismatch = pg.api.command(
        f"/api/v1/statements/{statement['id']}/approve",
        {
            "expected_version": saved["version"],
            **source,
            "sent_id": old_sent["id"],
            "feedback_id": feedback["id"],
            "quantity": "90",
            "unit_price": "10",
            "target_start": "2026-10-01",
            "target_end": "2026-10-31",
            "reason": "Synthetic mismatch",
        },
    )
    assert (
        mismatch.status_code == 409
        and mismatch.json()["code"] == "STATEMENT_EXTERNAL_NUMBER_MISMATCH"
    )
    export, sent = send_current(pg, statement)
    before = pg.api.request("GET", f"/api/v1/statement-exports/{export['id']}/file").content
    approval = approve(pg, statement, source, sent, feedback)
    assert approval["external_document_number"] == "SYN-EXT-001"
    review = command(
        pg,
        f"/api/v1/statements/{statement['id']}/review",
        {
            "expected_version": detail(pg, statement["id"])["version"],
            **source,
            "quantity": "10",
            "reason": "Synthetic review remainder",
        },
    )
    assert review["external_document_number"] == "SYN-EXT-001"
    reserve(pg, statement, approval, "60")
    formal = settle(pg, statement)
    assert formal["settlement_lines"][0]["external_document_number"] == "SYN-EXT-001"
    formal_export, _ = send_current(pg, statement)
    formal_bytes = pg.api.request(
        "GET", f"/api/v1/statement-exports/{formal_export['id']}/file"
    ).content
    later = blank(pg, "2026-11-01", "2026-11-30")
    claimed = command(
        pg,
        f"/api/v1/statements/{later['id']}/claim-review",
        {
            "expected_version": later["version"],
            "review_id": review["id"],
            "review_version": review["version"],
            "quantity": "5",
            "unit_price": "10",
            "reason": "Synthetic cross period",
        },
    )
    assert claimed["lines"][0]["external_document_number"] == "SYN-EXT-001"
    assert save_number(pg, later, "DIFFERENT").status_code == 409
    with pg.owner.begin() as conn:
        conn.execute(text("SELECT 1"))
        for table, identifier in (
            ("statement_approvals", approval["id"]),
            ("statement_reviews", review["id"]),
        ):
            with pytest.raises(DBAPIError), conn.begin_nested():
                conn.execute(
                    text(
                        f"UPDATE {table} SET external_document_number='OTHER',"
                        "version=version+1 WHERE id=:id"
                    ),
                    {"id": identifier},
                )
        with pytest.raises(DBAPIError), conn.begin_nested():
            conn.execute(
                text(
                    "UPDATE statement_lines SET external_document_number='OTHER' "
                    "WHERE statement_id=:id"
                ),
                {"id": statement["id"]},
            )
    assert (
        pg.api.request(
            "GET", f"/api/v1/statement-exports/{export['id']}/file?historical=true"
        ).content
        == before
    )
    assert (
        pg.api.request("GET", f"/api/v1/statement-exports/{formal_export['id']}/file").content
        == formal_bytes
    )


def test_pg_number_changes_participate_in_idempotency(pg):
    from test_postgres_statement_lifecycle import setup_statement

    statement, _, _, _ = setup_statement(pg)
    key = uuid4()
    first = save_number(pg, statement, "SYN-001", key=key)
    assert first.status_code == 200, first.text
    second = save_number(pg, statement, "SYN-002", key=key, expected=statement["version"])
    assert second.status_code == 409 and second.json()["code"] == "IDEMPOTENCY_CONFLICT"


def test_customer_excel_maps_external_number_as_text():
    from test_statement_customer_templates import custom_snapshot, template_file

    data = custom_snapshot()
    columns = data["customer_template"]["mapping"]["columns"]
    columns["external_document_number"] = columns.pop("item_code")
    data["statement"]["lines"][0]["external_document_number"] = "=SYN-OUT()"
    book = load_workbook(BytesIO(render_excel(data, template_file())), data_only=False)
    assert book.active["B4"].value == "=SYN-OUT()"
    assert book.active["B4"].data_type == "s"


def test_pg_review_claim_refuses_merge_with_different_number(pg):
    from test_postgres_statement_lifecycle import blank, command, detail, setup_statement

    statement, source, _, _ = setup_statement(pg)
    reviews = []
    for number in ("SYN-A", "SYN-B"):
        response = save_number(pg, statement, number)
        assert response.status_code == 200, response.text
        reviews.append(
            command(
                pg,
                f"/api/v1/statements/{statement['id']}/review",
                {
                    "expected_version": detail(pg, statement["id"])["version"],
                    **source,
                    "quantity": "5",
                    "reason": "Synthetic review identity",
                },
            )
        )
    target = blank(pg, "2026-11-01", "2026-11-30")
    for index, review in enumerate(reviews):
        response = pg.api.command(
            f"/api/v1/statements/{target['id']}/claim-review",
            {
                "expected_version": detail(pg, target["id"])["version"],
                "review_id": review["id"],
                "review_version": review["version"],
                "quantity": "5",
                "unit_price": "10",
                "reason": "Synthetic merge check",
            },
        )
        assert response.status_code == (200 if index == 0 else 409), response.text
    current = detail(pg, target["id"])
    assert current["lines"][0]["quantity"] == "5"
    assert current["lines"][0]["external_document_number"] == "SYN-A"
