"""Bounded synthetic import files; no runtime database or network access."""

import base64
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from factoryflow.kernel.contracts import BusinessError
from openpyxl import Workbook


def encoded(content):
    return base64.b64encode(content).decode()


def workbook_bytes(rows, configure=None):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    if configure:
        configure(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def read(content, filename="synthetic.csv", header_row=1):
    from factoryflow.modules.materials.adapters.import_file import read_import_file

    return read_import_file(filename, encoded(content), header_row)


def test_csv_preserves_text_codes_and_quoted_record_start_lines():
    table = read(
        '\ufeff我方料号,名称,规格\r\n000123,"合成\n支架",S-1\r\nA- 02,合成垫片,S-2\r\n'.encode()
    )
    assert table["rows"][0]["cells"][0]["text"] == "000123"
    assert table["rows"][0]["cells"][1]["text"] == "合成\n支架"
    assert [row["row_number"] for row in table["rows"]] == [2, 4]
    assert table["rows"][1]["cells"][0]["text"] == "A- 02"


def test_xlsx_preserves_text_and_reports_actual_numeric_value_and_format():
    raw = workbook_bytes(
        [["编号", "名称"], ["000123", "Synthetic text"], [123, "Synthetic numeric"]],
        lambda workbook: setattr(workbook.active["A3"], "number_format", "000000"),
    )
    table = read(raw, "synthetic.xlsx")
    assert table["rows"][0]["cells"][0]["text"] == "000123"
    numeric = table["rows"][1]["cells"][0]
    assert numeric["text"] == "123"
    assert numeric["kind"] == "number"
    assert numeric["number_format"] == "000000"


def test_xlsx_rejects_formula_after_inspection_sample():
    def formula(workbook):
        workbook.active["A100"] = "=1+1"

    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_UNSUPPORTED"):
        read(workbook_bytes([["名称"], ["Synthetic"]], formula), "synthetic.xlsx")


@pytest.mark.parametrize("feature", ["merge", "hidden", "hyperlink", "two_sheets"])
def test_xlsx_refuses_layouts_that_could_hide_or_combine_data(feature):
    def configure(workbook):
        sheet = workbook.active
        if feature == "merge":
            sheet.merge_cells("A2:B2")
        elif feature == "hidden":
            sheet.row_dimensions[2].hidden = True
        elif feature == "hyperlink":
            sheet["A2"].hyperlink = "https://synthetic.invalid/"
        else:
            workbook.create_sheet("Other").append(["Hidden synthetic data"])

    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_UNSUPPORTED"):
        read(workbook_bytes([["名称"], ["Synthetic"]], configure), "synthetic.xlsx")


def test_xlsx_allows_empty_spare_sheet():
    raw = workbook_bytes([["名称"], ["Synthetic"]], lambda workbook: workbook.create_sheet("Spare"))
    assert len(read(raw, "synthetic.xlsx")["rows"]) == 1


@pytest.mark.parametrize(
    "raw,filename",
    [
        (b"name\n\xff", "synthetic.csv"),
        (b'name\n"unterminated', "synthetic.csv"),
        (b"binary", "old.xls"),
    ],
)
def test_invalid_encoding_structure_and_unsupported_format_are_clear_errors(raw, filename):
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_(INVALID_FILE|UNSUPPORTED)"):
        read(raw, filename)


@pytest.mark.parametrize(
    "raw",
    [
        ("name\n" + "Synthetic\n" * 2001).encode(),
        (",".join(["name"] * 31) + "\nSynthetic").encode(),
        b"x" * (2 * 1024 * 1024 + 1),
    ],
    ids=["rows", "columns", "file_bytes"],
)
def test_whole_file_limits_do_not_silently_truncate(raw):
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_LIMIT"):
        read(raw)


def test_xlsx_uncompressed_limit_checked_before_loading_workbook():
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_LIMIT"):
        read(output.getvalue(), "synthetic.xlsx")


def test_header_row_is_explicit_and_bounded():
    table = read(b"Synthetic title\ncode,name\n000123,Synthetic\n", header_row=2)
    assert table["headers"] == ["code", "name"]
    assert table["rows"][0]["row_number"] == 3
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_MAPPING"):
        read(b"code,name\n000123,Synthetic\n", header_row=51)


def rewrite_sheet(raw, transform):
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = transform(data)
            target.writestr(entry.filename, data)
    return output.getvalue()


def test_malformed_cell_coordinates_are_a_safe_file_error():
    raw = workbook_bytes([["name"], ["Synthetic"]])
    altered = rewrite_sheet(raw, lambda data: data.replace(b'r="A2"', b'r="A0"'))
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_INVALID_FILE"):
        read(altered, "synthetic.xlsx")


def test_cached_dimension_cannot_hide_cells_outside_column_limit():
    raw = workbook_bytes([["name"], ["Synthetic"]])
    altered = rewrite_sheet(raw, lambda data: data.replace(b'r="A2"', b'r="AE2"'))
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_LIMIT"):
        read(altered, "synthetic.xlsx")


def test_shared_numeric_format_cannot_amplify_each_import_row():
    def large_shared_format(workbook):
        workbook.active["A2"].number_format = '"' + "x" * (1024 * 1024) + '"'
        workbook.active["A3"].number_format = workbook.active["A2"].number_format

    raw = workbook_bytes(
        [["own", "name"], [123, "Synthetic first"], [456, "Synthetic second"]], large_shared_format
    )
    with pytest.raises(BusinessError, match="MATERIAL_IMPORT_LIMIT"):
        read(raw, "synthetic.xlsx")
