from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier
from uuid import UUID, uuid4

import pytest
from factoryflow.kernel.contracts import Actor, BusinessError
from factoryflow.modules.statements.adapters.export_repository import StatementExportRepository
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from test_postgres_statements import draft, source


def create(pg, missing=False):
    delivery = source(pg)
    body = draft(pg, delivery)
    if missing:
        body["lines"][0]["unit_price"] = None
    statement = pg.api.command("/api/v1/statements", body).json()
    response = pg.api.command(
        f"/api/v1/statements/{statement['id']}/exports", {"expected_version": 1}
    )
    assert response.status_code == 200, response.text
    return delivery, body, statement, response.json()


def test_export_retry_failure_and_first_success_are_immutable(pg, monkeypatch):
    _, _, statement, export = create(pg, True)
    service = pg.runtime.statement_exports
    renderer = service.render

    def fail(snapshot):
        raise RuntimeError("synthetic renderer failure")

    monkeypatch.setattr(service, "render", fail)
    assert service.generate(pg.actor, export["id"])["status"] == "failed"
    monkeypatch.setattr(service, "render", renderer)
    ready = service.generate(pg.actor, export["id"])
    assert ready["status"] == "ready" and ready["current"]
    first, _ = service.download(pg.actor, export["id"])
    monkeypatch.setattr(service, "render", fail)
    assert service.generate(pg.actor, export["id"])["file_hash"] == ready["file_hash"]
    assert service.download(pg.actor, export["id"])[0] == first
    book = load_workbook(BytesIO(first))
    assert any(c.value == "待定价" for row in book.active for c in row)
    with pg.owner.connect() as conn:
        hashes = StatementExportRepository(conn).backup_hashes()
    assert len(hashes) == 1
    assert (pg.objects / f"{hashes[0]}.json").is_file()
    response = pg.api.request("GET", f"/api/v1/statement-exports/{export['id']}/file")
    assert response.status_code == 200 and response.content == first
    assert "no-store" in response.headers["cache-control"]


def test_export_locale_is_frozen_replayed_and_keeps_legacy_chinese_fingerprint(pg):
    _, _, statement, _ = create(pg, True)
    path = f"/api/v1/statements/{statement['id']}/exports"
    key = uuid4()
    old = pg.api.command(path, {"expected_version": 1}, key).json()
    explicit = pg.api.command(path, {"expected_version": 1, "locale": "zh-CN"}, key)
    assert explicit.status_code == 200, explicit.text
    assert explicit.json()["id"] == old["id"]
    assert explicit.json()["locale"] == "zh-CN"
    conflict = pg.api.command(path, {"expected_version": 1, "locale": "en"}, key)
    assert conflict.status_code == 409
    english_key = uuid4()
    body = {"expected_version": 1, "locale": "en"}
    english = pg.api.command(path, body, english_key)
    assert english.status_code == 200, english.text
    export = english.json()
    assert export["locale"] == "en"
    assert pg.api.command(path, body, english_key).json()["id"] == export["id"]
    service = pg.runtime.statement_exports
    service.generate(pg.actor, old["id"])
    old_bytes = service.download(pg.actor, old["id"])[0]
    service.generate(pg.actor, export["id"])
    en_bytes, metadata = service.download(pg.actor, export["id"])
    assert metadata["locale"] == "en"
    assert load_workbook(BytesIO(en_bytes)).active.title == "Statement Draft"
    assert service.download(pg.actor, old["id"])[0] == old_bytes
    assert pg.api.command(path, {"expected_version": 1, "locale": "fr"}).status_code == 422
    with pytest.raises(BusinessError) as failure:
        service.create(pg.actor, uuid4(), statement["id"], 1, locale="fr")
    assert failure.value.code == "UNSUPPORTED_LOCALE"


def test_export_replay_conflict_and_changed_draft_stay_historical(pg):
    _, body, statement, _ = create(pg)
    path = f"/api/v1/statements/{statement['id']}/exports"
    key = uuid4()
    response = pg.api.command(path, {"expected_version": 1}, key)
    export = response.json()
    saved = pg.api.command(
        f"/api/v1/statements/{statement['id']}/save", body | {"expected_version": 1}
    ).json()
    assert saved["version"] == 2
    replay = pg.api.command(path, {"expected_version": 1}, key).json()
    assert replay["id"] == export["id"] and not replay["current"]
    assert pg.api.command(path, {"expected_version": 2}, key).status_code == 409
    result = pg.api.request("POST", f"/api/v1/statement-exports/{export['id']}/generate")
    assert result.status_code == 200 and not result.json()["current"]
    assert (
        pg.api.request("GET", f"/api/v1/statement-exports/{export['id']}/file").status_code == 409
    )
    assert (
        pg.api.request(
            "GET", f"/api/v1/statement-exports/{export['id']}/file?historical=true"
        ).status_code
        == 200
    )


def test_source_void_during_render_cannot_reactivate_export(pg, monkeypatch):
    delivery, _, statement, export = create(pg)
    service = pg.runtime.statement_exports
    renderer = service.render

    def render(snapshot):
        result = pg.api.command(
            f"/api/v1/deliveries/{delivery['id']}/void",
            {"expected_version": 1, "reason": "synthetic correction"},
        )
        assert result.status_code == 200, result.text
        return renderer(snapshot)

    monkeypatch.setattr(service, "render", render)
    result = service.generate(pg.actor, export["id"])
    assert result["status"] == "ready" and not result["current"]
    current = pg.runtime.statements.detail(pg.actor, UUID(statement["id"]))
    assert current["calculation"]["total"] == "0.00"


def test_exported_draft_requires_explicit_void_and_releases_claim(pg):
    _, _, statement, export = create(pg)
    base = f"/api/v1/statements/{statement['id']}"
    body = {"expected_version": 1, "reason": "synthetic cancellation"}
    assert pg.api.command(base + "/abandon", body).json()["code"] == "STATEMENT_EXPORTED_USE_VOID"
    key = uuid4()
    response = pg.api.command(base + "/void", body, key)
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "void"
    assert pg.api.command(base + "/void", body, key).json() == response.json()
    assert not pg.runtime.statement_exports.generate(pg.actor, export["id"])["current"]
    with pg.owner.connect() as conn:
        assert conn.execute(text("SELECT sum(claimed) FROM statement_balances")).scalar() == 0


def test_concurrent_renderers_keep_first_bytes_and_late_failure_cannot_downgrade(pg, monkeypatch):
    _, _, _, export = create(pg)
    service = pg.runtime.statement_exports
    renderer, barrier = service.render, Barrier(2)

    def render(snapshot):
        barrier.wait()
        return renderer(snapshot)

    monkeypatch.setattr(service, "render", render)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: service.generate(pg.actor, export["id"]), range(2)))
    assert results[0]["file_hash"] == results[1]["file_hash"]
    with pg.runtime.engine.begin() as conn:
        StatementExportRepository(conn).finish(pg.actor, export["id"])
    assert service.generate(pg.actor, export["id"])["status"] == "ready"
    for query in [
        "UPDATE statement_exports SET snapshot='{}'",
        "UPDATE statement_exports SET file_hash=repeat('0',64)",
        "DELETE FROM statement_exports",
    ]:
        with pytest.raises(DBAPIError), pg.owner.begin() as conn:
            conn.execute(text(query))


def test_export_ownership_permissions_and_invalid_snapshot_rollback(pg):
    _, body, statement, export = create(pg)
    other = Actor(uuid4(), pg.actor.user_id, "synthetic-other", "operator")
    with pytest.raises(BusinessError) as failure:
        pg.runtime.statement_exports.download(other, export["id"])
    assert failure.value.code == "STATEMENT_EXPORT_NOT_FOUND"
    viewer = Actor(pg.actor.organization_id, pg.actor.user_id, "synthetic-viewer", "viewer")
    with pytest.raises(BusinessError):
        pg.runtime.statement_exports.generate(viewer, export["id"])
    body.update(
        reason="synthetic adjustment",
        adjustments=[
            {"stage": "after_tax", "kind": "deduction", "amount": "-100", "reason": "synthetic"}
        ],
    )
    assert (
        pg.api.command(
            f"/api/v1/statements/{statement['id']}/save", body | {"expected_version": 1}
        ).status_code
        == 200
    )
    path = f"/api/v1/statements/{statement['id']}/exports"
    assert (
        pg.api.command(path, {"expected_version": 2}).json()["code"]
        == "STATEMENT_EXPORT_REVIEW_REQUIRED"
    )
    with pg.owner.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM statement_exports")).scalar() == 1


def test_customer_template_default_and_old_snapshot_remain_independent(pg):
    _, _, statement, original = create(pg)
    pg.runtime.statements.settings.save(
        pg.actor,
        uuid4(),
        UUID(pg.customer_id),
        0,
        {
            "cycle_day": 25,
            "tax_mode": "add_tax",
            "tax_rate": "13",
            "template_id": "delivery-v1",
            "reason": "Synthetic preference",
        },
    )
    export = pg.runtime.statement_exports.create(pg.actor, uuid4(), statement["id"], 1)
    assert export["template_id"] == "delivery-v1"
    assert original["template_id"] == "standard-v1"
    pg.runtime.statement_exports.generate(pg.actor, export["id"])
    content, _ = pg.runtime.statement_exports.download(pg.actor, export["id"])
    book = load_workbook(BytesIO(content))
    assert any(str(c.value).startswith("送货：") for row in book.active for c in row)
    assert pg.runtime.statements.detail(pg.actor, statement["id"])["tax_mode"] == "untaxed_final"
