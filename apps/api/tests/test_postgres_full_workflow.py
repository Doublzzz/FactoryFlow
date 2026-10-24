"""One synthetic HTTP/PG workflow; restore only into a newly owned UUID database."""

import hashlib
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import psycopg
from conftest import API, ROOT
from factoryflow.bootstrap import create_app
from factoryflow.platform.backup import restore_backup, validate_backup
from factoryflow.platform.config import Settings
from openpyxl import load_workbook
from psycopg import sql
from sqlalchemy import text


def ok(response):
    assert response.status_code == 200, response.text
    return response.json()


def get_order(api, identifier):
    return ok(api.request("GET", f"/api/v1/orders/{identifier}"))


def test_structured_order_delivery_statement_export_void_and_restore(pg, tmp_path):
    line = {
        "item_code": "SYN-WORKFLOW-01",
        "item_name": "Synthetic cable",
        "specification": "Synthetic blue\nL=140mm, 4PIN",
        "quantity": "12.500",
        "unit": "KG",
        "source_unit_price": "0.08036",
        "source_amount": "0.99",
    }
    body = {
        "customer_id": pg.customer_id,
        "order_number": "SYN-FULL-WORKFLOW-01",
        "business_date": "2026-10-01",
        "currency": "CNY",
        "lines": [line],
    }
    order_key = uuid4()
    confirmed = ok(pg.api.command("/api/v1/orders", body, order_key))
    assert ok(pg.api.command("/api/v1/orders", body, order_key)) == confirmed
    order_id, revision_id = confirmed["order_id"], confirmed["revision_id"]
    original_order = get_order(pg.api, order_id)
    source_line = original_order["lines"][0]
    assert source_line["quantity"] == "12.5"
    assert source_line["source_quantity_text"] == "12.500"
    for field in ("item_code", "item_name", "specification", "source_unit_price", "source_amount"):
        assert source_line[field] == line[field]
    evidence_path = f"/api/v1/documents/{revision_id}/download"
    evidence = pg.api.request("GET", evidence_path)
    assert evidence.status_code == 200
    evidence_bytes = evidence.content
    assert hashlib.sha256(evidence_bytes).hexdigest() == confirmed["sha256"]
    assert evidence.json()["organization_id"] == str(pg.actor.organization_id)
    assert evidence.json()["actor_id"] == str(pg.actor.user_id)

    delivery_body = {
        "order_id": order_id,
        "customer_id": pg.customer_id,
        "revision_id": revision_id,
        "expected_version": 1,
        "delivery_date": "2026-10-02",
        "lines": [{"order_line_id": source_line["id"], "quantity": "4.125"}],
    }
    delivery_key = uuid4()
    delivery = ok(pg.api.command("/api/v1/deliveries", delivery_body, delivery_key))
    assert ok(pg.api.command("/api/v1/deliveries", delivery_body, delivery_key)) == delivery
    partial = get_order(pg.api, order_id)
    assert partial["version"] == 2
    assert partial["lines"][0]["delivered_quantity"] == "4.125"
    assert partial["lines"][0]["remaining_quantity"] == "8.375"

    layout_path = "/api/v1/delivery-layout"
    initial_layout = ok(pg.api.request("GET", layout_path))
    config = initial_layout["config"] | {"company_name": "Synthetic workflow factory"}
    layout = ok(pg.api.command(layout_path, {"expected_version": 0, "config": config}))
    assert layout["version"] == 1 and layout["config"] == config
    print_path = f"/api/v1/deliveries/{delivery['id']}/print"
    printed = pg.api.request("GET", print_path)
    assert printed.status_code == 200
    for value in ("Synthetic workflow factory", "SYN-WORKFLOW-01", "4.125", revision_id, "v1"):
        assert value in printed.text
    assert "0.08036" not in printed.text  # Delivery printing does not disclose source prices.
    assert get_order(pg.api, order_id) == partial

    statement_body = {
        "customer_id": pg.customer_id,
        "period_start": "2026-10-01",
        "period_end": "2026-10-31",
        "tax_mode": "add_tax",
        "tax_rate": "13",
        "adjustments": [],
        "lines": [
            {
                "delivery_id": delivery["id"],
                "order_line_id": source_line["id"],
                "quantity": "2.125",
                "unit_price": "2.3456",
                "reason": "Synthetic partial claim",
            }
        ],
    }
    statement = ok(pg.api.command("/api/v1/statements", statement_body))
    assert statement["lines"][0]["source"]["source_unit_price"] == "0.08036"
    assert statement["lines"][0]["source"]["source_amount"] == "0.99"
    assert statement["calculation"] == {
        "line_amounts": ["4.98"],
        "priced_subtotal": "4.98",
        "tax_base": "4.98",
        "tax_amount": "0.65",
        "total": "5.63",
        "warnings": [],
    }
    candidates = ok(
        pg.api.request(
            "GET",
            "/api/v1/statement-candidates",
            params={
                "customer_id": pg.customer_id,
                "period_start": "2026-10-01",
                "period_end": "2026-10-31",
            },
        )
    )
    assert len(candidates["items"]) == 1
    assert candidates["items"][0]["available_quantity"] == "2"
    export_path = f"/api/v1/statements/{statement['id']}/exports"
    export = ok(pg.api.command(export_path, {"expected_version": 1}))
    export_base = f"/api/v1/statement-exports/{export['id']}"
    ready = ok(pg.api.request("POST", export_base + "/generate"))
    assert ready["status"] == "ready" and ready["current"]
    downloaded = pg.api.request("GET", export_base + "/file")
    assert downloaded.status_code == 200
    export_bytes = downloaded.content
    assert hashlib.sha256(export_bytes).hexdigest() == ready["file_hash"]
    workbook = load_workbook(BytesIO(export_bytes))
    sheet = workbook.active
    assert [sheet.cell(7, col).value for col in range(2, 7)] == [
        "2026-10-02",
        line["item_code"],
        line["item_name"],
        line["specification"],
        "KG",
    ]
    assert [Decimal(str(sheet.cell(7, col).value)) for col in (7, 8, 9)] == [
        Decimal("2.125"),
        Decimal("2.3456"),
        Decimal("4.98"),
    ]
    assert [Decimal(str(sheet.cell(row, 10).value)) for row in range(8, 12)] == [
        Decimal("4.98"),
        Decimal("4.98"),
        Decimal("0.65"),
        Decimal("5.63"),
    ]
    assert sheet["J7"].value == "Synthetic partial claim"
    assert export["id"] in sheet["A13"].value
    workbook.close()

    void_body = {"expected_version": 1, "reason": "Synthetic workflow correction"}
    void_path = f"/api/v1/deliveries/{delivery['id']}/void"
    void_key = uuid4()
    voided = ok(pg.api.command(void_path, void_body, void_key))
    assert ok(pg.api.command(void_path, void_body, void_key)) == voided
    repeated = pg.api.command(void_path, void_body)
    assert repeated.status_code == 409 and repeated.json()["code"] == "DELIVERY_NOT_ACTIVE"
    reversed_order = get_order(pg.api, order_id)
    assert reversed_order["version"] == 3
    assert reversed_order["lines"][0]["remaining_quantity"] == "12.5"
    assert reversed_order["lines"][0]["delivered_quantity"] == "0"
    changed = ok(pg.api.request("GET", f"/api/v1/statements/{statement['id']}"))
    assert changed["version"] == 2
    assert changed["lines"][0]["status"] == "source_void"
    assert changed["calculation"]["total"] == "0.00"
    assert len(changed["history"]) == 2
    assert changed["history"][0]["snapshot"]["calculation"]["total"] == "5.63"
    stale = pg.api.request("GET", export_base + "/file")
    assert stale.status_code == 409 and stale.json()["code"] == "STATEMENT_EXPORT_STALE"
    historical = pg.api.request("GET", export_base + "/file?historical=true")
    assert historical.status_code == 200 and historical.content == export_bytes
    exports = ok(pg.api.request("GET", export_path))
    assert len(exports["items"]) == 1 and not exports["items"][0]["current"]
    assert exports["items"][0]["file_hash"] == ready["file_hash"]
    assert pg.api.request("GET", evidence_path).content == evidence_bytes
    with pg.owner.connect() as conn:
        assert conn.execute(text("SELECT sum(claimed) FROM statement_balances")).scalar() == 0

    backup = tmp_path / "workflow-backup"
    binary = ROOT / ".tools/postgresql-18.6-4/pgsql/bin"
    manifest = pg.runtime.backup(backup, binary)
    assert validate_backup(backup) == manifest
    assert confirmed["sha256"] in manifest["objects"]
    target_name = "ff_workflow_restore_" + uuid4().hex
    admin = pg.admin_url.set(drivername="postgresql", database="postgres")
    target_app = None
    with psycopg.connect(admin.render_as_string(hide_password=False), autocommit=True) as conn:
        assert (
            conn.execute("SELECT 1 FROM pg_database WHERE datname=%s", (target_name,)).fetchone()
            is None
        )
    try:
        objects = tmp_path / "restored-workflow-objects"
        restore_backup(admin, target_name, objects, backup, binary, ROOT)
        target_app = create_app(
            settings=Settings(
                pg.admin_url.set(database=target_name).render_as_string(hide_password=False),
                objects,
            )
        )
        restored_api = API(target_app)
        restored_api.login()  # Restore deliberately clears sessions; the original user survives.
        assert get_order(restored_api, order_id) == reversed_order
        assert ok(restored_api.request("GET", f"/api/v1/deliveries/{delivery['id']}")) == voided
        assert ok(restored_api.request("GET", f"/api/v1/statements/{statement['id']}")) == changed
        assert ok(restored_api.request("GET", layout_path)) == layout
        restored_evidence = restored_api.request("GET", evidence_path)
        assert restored_evidence.status_code == 200 and restored_evidence.content == evidence_bytes
        restored_export = restored_api.request("GET", export_base + "/file?historical=true")
        assert restored_export.status_code == 200 and restored_export.content == export_bytes
        assert restored_api.request("GET", export_base + "/file").status_code == 409
        assert ok(restored_api.request("GET", export_path)) == exports
        restored_print = restored_api.request("GET", print_path)
        assert restored_print.status_code == 200
        assert "Synthetic workflow factory" in restored_print.text
        assert "Synthetic workflow correction" in restored_print.text
        with target_app.state.services.engine.connect() as conn:
            assert conn.execute(text("SELECT sum(claimed) FROM statement_balances")).scalar() == 0
        for digest in manifest["objects"]:
            assert hashlib.sha256((objects / f"{digest}.json").read_bytes()).hexdigest() == digest
    finally:
        if target_app is not None:
            target_app.state.services.close()
        # This exact randomly named target was verified absent before this test's restore.
        assert target_name.startswith("ff_workflow_restore_") and len(target_name) == 52
        with psycopg.connect(admin.render_as_string(hide_password=False), autocommit=True) as conn:
            conn.execute(
                sql.SQL("DROP DATABASE IF EXISTS {} WITH (FORCE)").format(
                    sql.Identifier(target_name)
                )
            )
