"""Catch implicit currency conversion, ambiguous OCR and changed legacy fingerprints."""

import pytest
from factoryflow.kernel.contracts import BusinessError
from factoryflow.modules.orders.domain.order import prepare_order
from factoryflow.modules.statements.domain.draft import prepare_draft
from test_order_contract import payload


def test_usd_confirmation_retains_explicit_currency_and_exact_source_prices():
    body = payload() | {"currency": "USD"}
    order = prepare_order(body)
    assert order.currency == "USD"
    assert order.lines[0].source_amount == "99.00"
    assert order.lines[0].quantity.canonical == "1.125"


@pytest.mark.parametrize("currency", ["EUR", "usd", "$", "", None])
def test_unsupported_or_ambiguous_currency_is_not_assumed(currency):
    with pytest.raises(BusinessError, match="CURRENCY_REQUIRED"):
        prepare_order(payload() | {"currency": currency})


def test_statement_default_cny_keeps_legacy_fingerprint_and_usd_is_distinct():
    body = {
        "customer_id": payload()["customer_id"],
        "period_start": "2026-10-01",
        "period_end": "2026-10-31",
        "tax_mode": "untaxed_final",
        "tax_rate": "0",
        "lines": [],
    }
    legacy = prepare_draft(body)
    assert "currency" not in legacy
    assert prepare_draft(body | {"currency": "CNY"}) == legacy
    usd = prepare_draft(body | {"currency": "USD"})
    assert usd["currency"] == "USD"
    assert usd != legacy
    with pytest.raises(BusinessError, match="CURRENCY_REQUIRED"):
        prepare_draft(body | {"currency": "EUR"})


def test_ocr_explicit_usd_and_conflicting_currency_need_human_confirmation():
    from factoryflow.modules.documents.domain.candidate import propose

    assert propose(["Currency: USD", "Synthetic part 30 PCS"])["currency"] == "USD"
    assert propose(["币种：USD"])["currency"] == "USD"
    assert propose(["Currency: CNY", "Currency: USD"])["currency"] == ""
    assert propose(["Total: $100.00"])["currency"] == ""


@pytest.mark.parametrize("locale,label", [("en", "Total (USD)"), ("zh-CN", "最终金额（USD）")])
def test_export_currency_comes_from_snapshot_and_is_independent_of_language(locale, label):
    from io import BytesIO

    from factoryflow.modules.statements.adapters.excel import render_excel
    from openpyxl import load_workbook
    from test_statement_excel import snapshot

    data = snapshot() | {"locale": locale}
    data["statement"]["currency"] = "USD"
    data["statement"]["calculation"]["total"] = "70.37"
    sheet = load_workbook(BytesIO(render_excel(data))).active
    values = [cell.value for row in sheet for cell in row]
    assert label in values
    assert "USD" in sheet["A3"].value
    assert not any("CNY" in str(value) for value in values)
    total = next(cell for row in sheet for cell in row if cell.value == label)
    assert str(sheet.cell(total.row, 10).value) == "70.37"
