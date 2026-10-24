"""Read one bounded, flat material file. No files, network or business writes."""

import base64
import binascii
import csv
import hashlib
import re
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from factoryflow.kernel.contracts import BusinessError
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException, InvalidFileException

MAX_FILE = 2 * 1024 * 1024
MAX_EXPANDED = 20 * 1024 * 1024
MAX_DATA_ROWS, MAX_COLUMNS, MAX_HEADER = 2000, 30, 50
_PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/[\w.-]+\.xml|"
    r"xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|theme/theme\d+\.xml|"
    r"_rels/workbook\.xml\.rels|worksheets/sheet\d+\.xml|"
    r"worksheets/_rels/sheet\d+\.xml\.rels))\Z"
)


def invalid(field="file", code="MATERIAL_IMPORT_UNSUPPORTED"):
    raise BusinessError(code, field=field)


def cell_view(value, kind="text", number_format=None):
    if value is None:
        text = ""
    elif isinstance(value, bool):
        text, kind = ("TRUE" if value else "FALSE"), "boolean"
    elif isinstance(value, (datetime, date, time)):
        text, kind = value.isoformat(), "date"
    elif isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if not number.is_finite():
            invalid("cell", "MATERIAL_IMPORT_INVALID_FILE")
        text, kind = format(number, "f"), "number"
    else:
        text = str(value)
    if len(text) > 4000:
        invalid("cell", "MATERIAL_IMPORT_LIMIT")
    if number_format is not None and len(number_format.encode()) > 256:
        invalid("number_format", "MATERIAL_IMPORT_LIMIT")
    return {"text": text, "kind": kind, "number_format": number_format}


def check_xlsx(content):
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 400 or len({entry.filename for entry in entries}) != len(entries):
            invalid("zip")
        if sum(entry.file_size for entry in entries) > MAX_EXPANDED:
            invalid("file", "MATERIAL_IMPORT_LIMIT")
        for entry in entries:
            if entry.flag_bits & 1:
                invalid("zip")
            if entry.is_dir():
                continue
            if not _PART.fullmatch(entry.filename):
                invalid("workbook_parts")
            raw = archive.read(entry)
            if b"\x00" in raw or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                invalid("xml")
            root = ElementTree.fromstring(raw)
            for node in root.iter():
                tag = node.tag.rsplit("}", 1)[-1]
                if (
                    tag == "Relationship"
                    and node.attrib.get("TargetMode", "Internal") != "Internal"
                ):
                    invalid("external_links")
                if tag in {"f", "mergeCell", "oleObject", "externalReference", "tablePart"}:
                    invalid(tag)
                if tag == "sheet" and node.attrib.get("state", "visible") != "visible":
                    invalid("hidden_sheet")
                if tag in {"row", "col"} and node.attrib.get("hidden", "0") in {"1", "true"}:
                    invalid("hidden_data")
                if tag == "filterColumn":
                    invalid("filtered_data")
                if entry.filename.startswith("xl/worksheets/"):
                    references = []
                    if tag == "c":
                        references = [node.attrib.get("r", "")]
                    elif tag == "dimension":
                        references = node.attrib.get("ref", "").split(":")
                    for reference in references:
                        column, row = coordinate_from_string(reference)
                        if (
                            row > MAX_DATA_ROWS + MAX_HEADER
                            or column_index_from_string(column) > MAX_COLUMNS
                        ):
                            invalid("rows_columns", "MATERIAL_IMPORT_LIMIT")


def xlsx_rows(content):
    check_xlsx(content)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        populated = []
        total_text = 0
        for sheet in workbook.worksheets:
            # XML coordinates were bounded independently; do not trust cached dimensions.
            sheet.reset_dimensions()
            rows = []
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if number > MAX_DATA_ROWS + MAX_HEADER or len(cells) > MAX_COLUMNS:
                    invalid("rows_columns", "MATERIAL_IMPORT_LIMIT")
                values = []
                for cell in cells:
                    kind = "error" if cell.data_type == "e" else "text"
                    value = cell_view(
                        cell.value, kind, cell.number_format if cell.value is not None else None
                    )
                    total_text += len(value["text"].encode()) + len(
                        (value["number_format"] or "").encode()
                    )
                    if total_text > MAX_EXPANDED:
                        invalid("cell_text", "MATERIAL_IMPORT_LIMIT")
                    values.append(value)
                if any(value["text"] for value in values):
                    rows.append({"row_number": number, "cells": values})
            if rows:
                populated.append((sheet.title, rows))
        if len(populated) != 1:
            invalid("sheets")
        return populated[0]
    finally:
        workbook.close()


def csv_rows(content):
    reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
    rows = []
    number = 0
    for index, values in enumerate(reader, 1):
        start = number + 1
        number = reader.line_num
        if index > MAX_DATA_ROWS + MAX_HEADER or len(values) > MAX_COLUMNS:
            invalid("rows_columns", "MATERIAL_IMPORT_LIMIT")
        if any(values):
            rows.append({"row_number": start, "cells": [cell_view(value) for value in values]})
    return "CSV", rows


def read_import_file(filename, content_base64, header_row=1):
    if type(header_row) is not int or not 1 <= header_row <= MAX_HEADER:
        invalid("header_row", "MATERIAL_IMPORT_MAPPING")
    try:
        if not isinstance(content_base64, str) or len(content_base64) > 2796204:
            invalid("file", "MATERIAL_IMPORT_LIMIT")
        content = base64.b64decode(content_base64, validate=True)
        if not content or len(content) > MAX_FILE:
            invalid("file", "MATERIAL_IMPORT_LIMIT")
        extension = PurePosixPath(filename).suffix.lower()
        if extension == ".csv":
            sheet_name, rows = csv_rows(content)
        elif extension == ".xlsx":
            sheet_name, rows = xlsx_rows(content)
        else:
            invalid("format")
    except BusinessError:
        raise
    except (
        ValueError,
        TypeError,
        KeyError,
        OSError,
        UnicodeError,
        csv.Error,
        BadZipFile,
        binascii.Error,
        ElementTree.ParseError,
        CellCoordinatesException,
        InvalidFileException,
    ):
        invalid("file", "MATERIAL_IMPORT_INVALID_FILE")
    header = next((row for row in rows if row["row_number"] == header_row), None)
    if header is None:
        invalid("header_row", "MATERIAL_IMPORT_MAPPING")
    data = [row for row in rows if row["row_number"] > header_row]
    if len(data) > MAX_DATA_ROWS:
        invalid("rows", "MATERIAL_IMPORT_LIMIT")
    width = max((len(row["cells"]) for row in rows), default=0)
    for row in data:
        row["cells"] += [cell_view(None) for _ in range(width - len(row["cells"]))]
    return {
        "file_hash": hashlib.sha256(content).hexdigest(),
        "sheet_name": sheet_name,
        "headers": [cell["text"] for cell in header["cells"]]
        + [""] * (width - len(header["cells"])),
        "header_row": header_row,
        "rows": data,
    }
