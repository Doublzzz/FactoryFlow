"""Fixed printable layouts; amounts and approvals come from the saved snapshot."""

from decimal import Decimal
from io import BytesIO
from math import ceil

from factoryflow.modules.statements.domain.templates import TEMPLATES
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def render_excel(snapshot, template_content=None):
    locale = snapshot.get("locale", "zh-CN")
    if locale not in ("zh-CN", "en"):
        raise ValueError("Unsupported export locale")

    def text(chinese, english):
        return english if locale == "en" else chinese

    if snapshot["template_id"] == "customer-v1":
        from factoryflow.modules.statements.adapters.customer_excel import render_customer_excel

        return render_customer_excel(snapshot, template_content)
    if snapshot["template_id"] not in TEMPLATES or snapshot["template_version"] != 1:
        raise ValueError("Unsupported immutable layout version")
    data = snapshot["statement"]
    formal = data.get("status", "draft") == "confirmed"
    label = text("正式", "Final") if formal else text("草稿", "Draft")
    status_label = label
    historical = snapshot.get("current") is False
    display_label = text(f"历史版本 · {label}", f"Historical {label}") if historical else label
    book = Workbook()
    sheet = book.active
    sheet.title = text(f"对账{display_label}", f"Statement {display_label}")
    sheet.sheet_view.showGridLines = False
    # Keep ten business columns readable within an A4 portrait page.
    widths = [5, 12, 9, 12, 16, 5, 9, 9, 11, 12]
    thin = Side(style="thin", color="D6DCE3")

    def cell(row, col, value, *, number=False, bold=False):
        target = sheet.cell(row, col)
        if value is not None:
            # Excel preserves at most 15 significant digits. Large values stay exact text.
            digits = len(Decimal(str(value)).as_tuple().digits) if number else 0
            if number and digits <= 15:
                target.value = Decimal(str(value))
                target.number_format = (
                    "0.00"
                    if col in (9, 10)
                    else "0"
                    if target.value == target.value.to_integral_value()
                    else "0.####"
                )
            else:
                target.value = str(value)
                target.data_type = "s"  # User text must never become an Excel formula.
        target.font = Font(name="Microsoft YaHei", size=10, bold=bold, color="223047")
        target.alignment = Alignment(
            vertical="center", wrap_text=True, horizontal="right" if number else "left"
        )
        target.border = Border(bottom=thin)
        return target

    def merged(row, start, end, value, bold=False):
        sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        available = sum(widths[start - 1 : end]) - 2
        lines = sum(
            max(1, ceil(sum(2 if ord(c) > 255 else 1 for c in part) / available))
            for part in str(value).split("\n")
        )
        sheet.row_dimensions[row].height = max(20, lines * 16)
        return cell(row, start, value, bold=bold)

    merged(
        1,
        1,
        10,
        text(f"客户对账单 · {display_label}", f"Reconciliation statement · {display_label}"),
        True,
    ).font = Font(name="Microsoft YaHei", size=16, bold=True)
    sheet.row_dimensions[1].height = 32
    merged(
        2, 1, 10, text(f"客户：{data['customer_name']}", f"Customer: {data['customer_name']}"), True
    )
    merged(
        3,
        1,
        10,
        text(
            f"期间：{data['period_start']} 至 {data['period_end']}    币种：CNY"
            f"    {label}版本：{data['version']}",
            f"Period: {data['period_start']} to {data['period_end']}    Currency: CNY"
            f"    {label} version: {data['version']}",
        ),
    )
    mode = {
        "included": text("含税价（不另加税）", "Tax included"),
        "untaxed_final": text("未税最终价（不另加税）", "No additional tax"),
        "add_tax": text("未税加税", "Add tax"),
    }[data["tax_mode"]]
    confirmation = (
        text("本文件保存正式确认时的固定版本", "Fixed version at final confirmation")
        if formal
        else text(
            "未经客户确认，不代表正式结算",
            "Not yet accepted by the customer; not a final statement",
        )
    )
    merged(
        4,
        1,
        10,
        text(
            f"计价方式：{mode}    税率：{data['tax_rate']}%    {confirmation}",
            f"Pricing: {mode}    Tax rate: {data['tax_rate']}%    {confirmation}",
        ),
    )
    headers = [
        "序号",
        "送货日期",
        "货物编号",
        "货物名称",
        "规格型号",
        "单位",
        "本次数量",
        "结算单价",
        "行金额",
        "说明",
    ]
    if locale == "en":
        headers = [
            "No.",
            "Delivery date",
            "Item code",
            "Item name",
            "Specification",
            "Unit",
            "Quantity",
            "Unit price",
            "Amount",
            "Notes",
        ]
    for index, label in enumerate(headers, 1):
        cell(6, index, label, bold=True).fill = PatternFill("solid", fgColor="EAF0F6")
    sheet.row_dimensions[5].height = 8
    sheet.row_dimensions[6].height = 36 if locale == "en" else 30
    row, ordinal, previous = 7, 0, None
    amounts = data["calculation"]["line_amounts"]
    all_lines = data["lines"] + data.get("settlement_lines", [])
    if len(amounts) != len(all_lines):
        raise ValueError("Saved calculation does not match statement details")
    lines = [(i, line) for i, line in enumerate(all_lines) if line["status"] == "included"]
    if snapshot["template_id"] == "delivery-v1":
        lines.sort(
            key=lambda pair: (pair[1]["source"]["delivery_date"], pair[1]["delivery_id"], pair[0])
        )
    for original, line in lines:
        source = line["source"]
        if snapshot["template_id"] == "delivery-v1" and previous != line["delivery_id"]:
            previous = line["delivery_id"]
            merged(
                row,
                1,
                10,
                text(
                    f"送货：{source['delivery_date']}  /  {previous}",
                    f"Delivery: {source['delivery_date']}  /  {previous}",
                ),
                True,
            )
            row += 1
        ordinal += 1
        values = [
            ordinal,
            source["delivery_date"],
            source.get("item_code"),
            source.get("item_name") or source.get("description"),
            source.get("specification"),
            source.get("canonical_unit") or source.get("source_unit"),
            line["quantity"],
            line["unit_price"]
            if line["unit_price"] is not None
            else text("待定价", "Price pending"),
            amounts[original],
            " / ".join(
                part
                for part in (
                    line["reason"],
                    text(
                        f"外部单号：{line['external_document_number']}",
                        f"External number: {line['external_document_number']}",
                    )
                    if line.get("external_document_number")
                    else "",
                    text(
                        f"认可：{line['approval_id']} v{line['approval_version']}",
                        f"Acceptance: {line['approval_id']} v{line['approval_version']}",
                    )
                    if "approval_id" in line
                    else "",
                    text(
                        f"分配：{line['allocation_id']}"
                        f"（{'已结算' if line.get('allocation_status') == 'settled' else '预留'}）",
                        f"Allocation: {line['allocation_id']}"
                        " ("
                        + (
                            "finalized"
                            if line.get("allocation_status") == "settled"
                            else "reserved"
                        )
                        + ")",
                    )
                    if "allocation_id" in line
                    else "",
                )
                if part
            ),
        ]
        for col, value in enumerate(values, 1):
            cell(
                row,
                col,
                value,
                number=col in (7, 8, 9)
                and value is not None
                and not (col == 8 and line["unit_price"] is None),
            )
        # Wrapped long specifications get room; no shortening of saved business text.
        height = max(
            sum(
                max(1, ceil(sum(2 if ord(c) > 255 else 1 for c in part) / widths[i]))
                for part in str(v or "").split("\n")
            )
            for i, v in enumerate(values)
        )
        sheet.row_dimensions[row].height = min(409, max(28, height * 16))
        row += 1
    for item in data["adjustments"]:
        label = (
            text("税前", "Before-tax")
            if item["stage"] == "before_tax"
            else text("税后", "After-tax")
        )
        merged(
            row,
            1,
            9,
            text(
                f"{label}调整 / {item['kind']}：{item['reason']}",
                f"{label} adjustment / {item['kind']}: {item['reason']}",
            ),
        )
        cell(row, 10, item["amount"], number=True)
        row += 1
    calculation = data["calculation"]
    for label, key in [
        (text("已定价小计", "Priced subtotal"), "priced_subtotal"),
        (text("计税基数", "Tax base"), "tax_base"),
        (text("税额", "Tax amount"), "tax_amount"),
        (text("最终金额（CNY）", "Total (CNY)"), "total"),
    ]:
        merged(row, 1, 9, label, True).alignment = Alignment(
            vertical="center", horizontal="right", wrap_text=True
        )
        cell(row, 10, calculation[key], number=True, bold=True)
        row += 1
    if "PRICE_REQUIRED" in calculation["warnings"]:
        merged(
            row,
            1,
            10,
            text(
                "存在待定价明细；小计仅含已定价部分，最终金额尚未确定。",
                "Some lines have no price. The subtotal includes priced lines only; "
                "the total is not yet determined.",
            ),
        )
        row += 1
    merged(
        row + 1,
        1,
        10,
        text(
            f"导出编号：{snapshot['id']}    模板：{TEMPLATES[snapshot['template_id']]} v1",
            f"Export ID: {snapshot['id']}    Layout: "
            f"{'Standard' if snapshot['template_id'] == 'standard-v1' else 'By delivery'} v1",
        ),
    )
    merged(
        row + 2,
        1,
        10,
        text(
            f"本文件保存导出时的{label}。历史文件字节保留；发送或结算前请在系统核对当前版本。",
            f"This file preserves the {status_label.lower()} snapshot. "
            "Check the current version before sending or finalizing reconciliation. "
            "It does not record payment.",
        ),
    )
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "G7"
    sheet.print_title_rows = "1:6"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_area = f"A1:J{row + 2}"
    sheet.oddFooter.center.text = text(
        f"第 &P 页 / 共 &N 页 · 对账{display_label}", f"Page &P of &N · Statement {display_label}"
    )
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
