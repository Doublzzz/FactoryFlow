"""One content sheet, explicit bindings, and dependency-free row rebuilding."""

import hashlib
import posixpath
import re
import warnings
from colorsys import hls_to_rgb, rgb_to_hls
from copy import copy
from decimal import Decimal
from io import BytesIO
from math import ceil
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from factoryflow.kernel.contracts import BusinessError
from factoryflow.modules.statements.domain.customer_templates import (
    CELL_FIELDS,
    DETAIL_FIELDS,
    REQUIRED_CELLS,
    REQUIRED_COLUMNS,
    require_coverage,
)
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.print_settings import PrintArea

MAX_FILE = 2 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS = 300, 30
_PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core|custom)\.xml|"
    r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/theme/theme1\.xml|"
    r"xl/_rels/workbook\.xml\.rels|xl/worksheets/sheet[0-9]+\.xml)\Z"
)
_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_SHEET_TAGS = set(
    (
        "worksheet sheetPr tabColor outlinePr pageSetUpPr dimension sheetViews sheetView "
        "pane selection sheetFormatPr cols col sheetData row c f v is t mergeCells mergeCell "
        "printOptions pageMargins pageSetup headerFooter oddHeader oddFooter evenHeader "
        "evenFooter firstHeader firstFooter autoFilter r rPr rFont charset family b i strike "
        "outline shadow condense extend color sz u vertAlign scheme"
    ).split()
)
_CALC_NS = "{http://schemas.microsoft.com/office/spreadsheetml/2018/calcfeatures}"
_WPS_FILTER = "{http://www.wps.cn/officeDocument/2017/etCustomData}filterBottomFollowUsedRange"


def _calculation_metadata(root):
    """Only understood capability flags may be discarded by the value-only writer."""
    for extensions in root.findall(_NS + "extLst"):
        if extensions.attrib or len(extensions) != 1:
            invalid("workbook_features")
        extension = extensions[0]
        if (
            extension.tag != _NS + "ext"
            or extension.attrib != {"uri": "{B58B0392-4F1F-4190-BB64-5DF3571DCE5F}"}
            or len(extension) != 1
        ):
            invalid("workbook_features")
        features = extension[0]
        if features.tag != _CALC_NS + "calcFeatures" or features.attrib:
            invalid("workbook_features")
        names = {"RD", "Single", "FV", "CNMTM", "LET_WF", "LAMBDA_WF", "ARRAYTEXT_WF"}
        for feature in features:
            if (
                feature.tag != _CALC_NS + "feature"
                or set(feature.attrib) != {"name"}
                or feature.attrib["name"] not in {"microsoft.com:" + name for name in names}
                or len(feature)
                or (feature.text or "").strip()
            ):
                invalid("workbook_features")
        root.remove(extensions)


def _empty_filters(root):
    for filter_node in root.findall(_NS + "autoFilter"):
        if set(filter_node.attrib) - {"ref", _WPS_FILTER}:
            invalid("filters")
        if _WPS_FILTER in filter_node.attrib and filter_node.attrib[_WPS_FILTER] != "0":
            invalid("filters")
        if filter_node.attrib.get("ref"):
            bounded_range(filter_node.attrib["ref"])
        for child in list(filter_node):
            if (
                child.tag != _NS + "extLst"
                or child.attrib
                or len(child)
                or (child.text or "").strip()
            ):
                invalid("filters")
            filter_node.remove(child)


def _string_properties(root):
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}"
    value_tag = "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}lpwstr"
    if root.tag != namespace + "Properties" or root.attrib:
        invalid("custom_properties")
    for prop in root:
        if (
            prop.tag != namespace + "property"
            or set(prop.attrib) != {"fmtid", "pid", "name"}
            or prop.attrib["fmtid"] != "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
            or len(prop) != 1
            or prop[0].tag != value_tag
            or prop[0].attrib
            or len(prop[0])
        ):
            invalid("custom_properties")


def _unused_slicer_style(root):
    for extensions in root.findall(_NS + "extLst"):
        if extensions.attrib or len(extensions) != 1:
            invalid("style_extensions")
        extension = extensions[0]
        if (
            extension.tag != _NS + "ext"
            or extension.attrib != {"uri": "{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}"}
            or len(extension) != 1
        ):
            invalid("style_extensions")
        style = extension[0]
        if (
            style.tag
            != "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}slicerStyles"
            or style.attrib != {"defaultSlicerStyle": "SlicerStyleLight1"}
            or len(style)
        ):
            invalid("style_extensions")
        root.remove(extensions)


def invalid(field=None, code="STATEMENT_TEMPLATE_UNSUPPORTED"):
    raise BusinessError(code, 422, field)


def bounded_range(value):
    try:
        lo_col, lo_row, hi_col, hi_row = range_boundaries(value)
        if not (1 <= lo_col <= hi_col <= MAX_COLUMNS and 1 <= lo_row <= hi_row <= MAX_ROWS):
            invalid("range")
        return lo_col, lo_row, hi_col, hi_row
    except TypeError, ValueError:
        invalid("range")


def load_template(content):
    if not content or len(content) > MAX_FILE:
        invalid("file", "STATEMENT_TEMPLATE_FILE_LIMIT")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {entry.filename for entry in infos}
            if len(infos) > 400 or len(names) != len(infos):
                invalid("zip")
            if sum(entry.file_size for entry in infos) > 20 * 1024 * 1024:
                invalid("file", "STATEMENT_TEMPLATE_FILE_LIMIT")
            directories = {
                name[: position + 1]
                for name in names
                if _PART.fullmatch(name)
                for position, char in enumerate(name)
                if char == "/"
            }
            for entry in infos:
                if entry.flag_bits & 1:
                    invalid("zip")
                if entry.is_dir():
                    if entry.filename not in directories or entry.file_size:
                        invalid("zip")
                    continue
                if not _PART.fullmatch(entry.filename):
                    invalid("zip")
                raw = archive.read(entry)
                if b"\x00" in raw or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                    invalid("xml")
                root = ElementTree.fromstring(raw)
                if entry.filename.endswith(".rels"):
                    for relation in root:
                        kind = relation.attrib.get("Type", "").rsplit("/", 1)[-1]
                        target = relation.attrib.get("Target", "")
                        base = "xl" if entry.filename.startswith("xl/") else ""
                        resolved = posixpath.normpath(posixpath.join(base, target.lstrip("/")))
                        if target.startswith("/"):
                            resolved = target.lstrip("/")
                        if (
                            relation.attrib.get("TargetMode", "Internal") != "Internal"
                            or kind
                            not in {
                                "officeDocument",
                                "core-properties",
                                "extended-properties",
                                "custom-properties",
                                "worksheet",
                                "styles",
                                "theme",
                                "sharedStrings",
                            }
                            or resolved not in names
                        ):
                            invalid("relationships")
                if entry.filename.startswith("xl/worksheets/"):
                    _empty_filters(root)
                    for node in root.iter():
                        if not node.tag.startswith(_NS) or node.tag[len(_NS) :] not in _SHEET_TAGS:
                            invalid("worksheet_features")
                        tag = node.tag[len(_NS) :]
                        if tag in {"dimension", "mergeCell"}:
                            bounded_range(node.attrib.get("ref", ""))
                        if tag == "c":
                            bounded_range(node.attrib.get("r", ""))
                        if tag == "row" and not 1 <= int(node.attrib.get("r", "0")) <= MAX_ROWS:
                            invalid("rows")
                        if tag == "row" and int(node.attrib.get("outlineLevel", "0")):
                            invalid("row_groups")
                        if tag == "col" and not (
                            1 <= int(node.attrib["min"]) <= int(node.attrib["max"]) <= MAX_COLUMNS
                        ):
                            invalid("columns")
                        if tag == "col" and int(node.attrib.get("outlineLevel", "0")):
                            invalid("column_groups")
                        if tag == "f" and node.attrib:
                            invalid("formula_kind")
                if entry.filename == "xl/workbook.xml":
                    _calculation_metadata(root)
                    allowed = {
                        "workbook",
                        "fileVersion",
                        "workbookPr",
                        "bookViews",
                        "workbookView",
                        "sheets",
                        "sheet",
                        "definedNames",
                        "definedName",
                        "calcPr",
                        "workbookProtection",
                    }
                    for node in root.iter():
                        if not node.tag.startswith(_NS) or node.tag[len(_NS) :] not in allowed:
                            invalid("workbook_features")
                        if node.tag == _NS + "workbookProtection" and node.attrib:
                            invalid("workbook_protection")
                        if (
                            node.tag == _NS + "definedName"
                            and node.attrib.get("name") == "_xlnm._FilterDatabase"
                        ):
                            if (
                                node.attrib.get("hidden") != "1"
                                or "localSheetId" not in node.attrib
                                or "!" not in (node.text or "")
                            ):
                                invalid("named_ranges")
                            bounded_range(node.text.rsplit("!", 1)[1])
                            continue
                        if node.tag == _NS + "definedName" and node.attrib.get("name") not in {
                            "_xlnm.Print_Area",
                            "_xlnm.Print_Titles",
                        }:
                            invalid("named_ranges")
                if entry.filename == "[Content_Types].xml":
                    if b"macroEnabled" in raw or b"vba" in raw.lower():
                        invalid("macros")
                if entry.filename == "docProps/custom.xml":
                    _string_properties(root)
                if entry.filename in {"xl/styles.xml", "xl/sharedStrings.xml"}:
                    if len(list(root.iter())) > 50000:
                        invalid("styles", "STATEMENT_TEMPLATE_FILE_LIMIT")
                    if entry.filename == "xl/styles.xml":
                        _unused_slicer_style(root)
                    if root.findall(".//" + _NS + "extLst"):
                        invalid("style_extensions")
                    if entry.filename == "xl/sharedStrings.xml":
                        allowed = (_SHEET_TAGS - {"worksheet"}) | {"sst", "si"}
                        if any(
                            not n.tag.startswith(_NS) or n.tag[len(_NS) :] not in allowed
                            for n in root.iter()
                        ):
                            invalid("rich_text_features")
                    palette = root.find(".//" + _NS + "indexedColors")
                    if palette is not None and [node.attrib.get("rgb") for node in palette] != list(
                        COLOR_INDEX
                    ):
                        invalid("custom_palette")
        with warnings.catch_warnings(record=True) as notices:
            warnings.simplefilter("always")
            book = load_workbook(
                BytesIO(content), data_only=False, keep_links=False, rich_text=True
            )
        populated = [
            s for s in book.worksheets if any(c.value is not None for row in s for c in row)
        ]
        if notices or len(populated) != 1 or populated[0].sheet_state != "visible":
            invalid("sheets")
        book.active = book.worksheets.index(populated[0])
        sheet = book.active
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
            invalid("dimensions", "STATEMENT_TEMPLATE_FILE_LIMIT")
        return book
    except BusinessError:
        raise
    except (
        BadZipFile,
        ValueError,
        KeyError,
        TypeError,
        OSError,
        ElementTree.ParseError,
        OverflowError,
    ):
        invalid("file")


def grid(sheet):
    return [
        [str(cell.value) if cell.value is not None else "" for cell in row]
        for row in sheet.iter_rows()
    ]


def preview_grid(content):
    return grid(load_workbook(BytesIO(content)).active)


def inspect_template(content):
    sheet = load_template(content).active
    aliases = {
        "ordinal": {"序号"},
        "delivery_date": {"日期", "送货日期"},
        "item_code": {"货号", "编号", "货物编号", "物料编码"},
        "item_name_specification": {"品名/规格", "名称/规格", "品名规格", "品名／规格"},
        "order_number": {"订单号", "采购单号"},
        "item_name": {"名称", "品名", "货物名称"},
        "specification": {"规格", "规格型号"},
        "unit": {"单位"},
        "quantity": {"数量", "本次数量"},
        "unit_price": {"单价", "结算单价"},
        "amount": {"金额", "行金额"},
        "reason": {"说明", "备注"},
    }
    candidates = []
    for row in sheet:
        found = {}
        for cell in row:
            for field, names in aliases.items():
                label = re.sub(r"\s+", "", str(cell.value or ""))
                if label in names and field not in found:
                    found[field] = get_column_letter(cell.column)
        if len(found) >= 3:
            candidates.append((len(found), row[0].row, found))
    suggested = None
    if candidates:
        _, header, columns = max(candidates, key=lambda item: (item[0], -item[1]))
        summaries = [
            cell.row
            for row in sheet
            for cell in row
            if cell.row > header
            and str(cell.value or "").strip() in {"合计", "总计", "已定价小计", "最终金额"}
        ]
        end = min(summaries) - 1 if summaries else min(header + 1, sheet.max_row)
        if header < end:
            suggested = {
                "sheet_name": sheet.title,
                "detail_start_row": header + 1,
                "detail_end_row": end,
                "columns": columns,
                "cells": {},
            }
    return {
        "file_hash": hashlib.sha256(content).hexdigest(),
        "sheet_name": sheet.title,
        "rows": sheet.max_row,
        "columns": sheet.max_column,
        "grid": grid(sheet),
        "suggested_mapping": suggested,
    }


def validate_mapping(content, mapping):
    book = load_template(content)
    sheet = book.active
    if set(mapping) != {"sheet_name", "detail_start_row", "detail_end_row", "columns", "cells"}:
        invalid("mapping", "STATEMENT_TEMPLATE_MAPPING")
    start, end = mapping["detail_start_row"], mapping["detail_end_row"]
    if (
        type(start) is not int
        or type(end) is not int
        or not 2 <= start <= end <= sheet.max_row
        or mapping["sheet_name"] != sheet.title
    ):
        invalid("detail_rows", "STATEMENT_TEMPLATE_MAPPING")
    columns, cells = mapping["columns"], mapping["cells"]
    if (
        not isinstance(columns, dict)
        or not isinstance(cells, dict)
        or not REQUIRED_COLUMNS <= columns.keys() <= DETAIL_FIELDS
        or not ({"item_name", "item_name_specification"} & columns.keys())
        or not REQUIRED_CELLS <= cells.keys() <= CELL_FIELDS
    ):
        invalid("required_fields", "STATEMENT_TEMPLATE_MAPPING")
    if any(not isinstance(v, str) or not re.fullmatch(r"[A-Z]{1,2}", v) for v in columns.values()):
        invalid("columns", "STATEMENT_TEMPLATE_MAPPING")
    if any(column_index_from_string(v) > sheet.max_column for v in columns.values()):
        invalid("columns", "STATEMENT_TEMPLATE_MAPPING")
    if any(
        not isinstance(v, str) or not re.fullmatch(r"[A-Z]{1,2}[1-9][0-9]{0,2}", v)
        for v in cells.values()
    ):
        invalid("cells", "STATEMENT_TEMPLATE_MAPPING")
    if len(set(columns.values())) != len(columns) or len(set(cells.values())) != len(cells):
        invalid("duplicate_bindings", "STATEMENT_TEMPLATE_MAPPING")
    targets = set(cells.values()) | {
        f"{col}{row}" for col in columns.values() for row in range(start, end + 1)
    }
    print_left, print_right = _print_columns(sheet)
    for address in targets:
        _, row, col, _ = bounded_range(address)
        if row > sheet.max_row or col > sheet.max_column:
            invalid(address, "STATEMENT_TEMPLATE_MAPPING")
        if not print_left <= col <= print_right:
            invalid(address, "STATEMENT_TEMPLATE_MAPPING")
        merged = next((area for area in sheet.merged_cells.ranges if address in area), None)
        if merged and (merged.min_col < print_left or merged.max_col > print_right):
            invalid(address, "STATEMENT_TEMPLATE_MAPPING")
        cell = sheet[address]
        hidden = any(d.hidden and d.min <= col <= d.max for d in sheet.column_dimensions.values())
        if isinstance(cell, MergedCell) or sheet.row_dimensions[row].hidden or hidden:
            invalid(address, "STATEMENT_TEMPLATE_MAPPING")
        _foreground(sheet, cell)
    for field, address in cells.items():
        row = sheet[address].row
        if (field in {"customer_name", "period", "status"} and row >= start) or (
            field not in {"customer_name", "period", "status"} and row <= end
        ):
            invalid(address, "STATEMENT_TEMPLATE_MAPPING")
    first_summary = min(sheet[cells[k]].row for k in {"priced_subtotal", "total"} if k in cells)
    if first_summary not in {end + 1, end + 2}:
        invalid("summary_gap", "STATEMENT_TEMPLATE_MAPPING")
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= end and merged.max_row >= start and merged.min_row != merged.max_row:
            invalid(str(merged), "STATEMENT_TEMPLATE_MAPPING")
    for row in sheet:
        for cell in row:
            if cell.data_type == "f" and cell.coordinate not in targets:
                invalid(cell.coordinate, "STATEMENT_TEMPLATE_FORMULA")
    if sheet.print_title_cols:
        invalid("print_title_cols")
    if sheet.print_title_rows:
        match = re.fullmatch(r"\$?([0-9]+):\$?([0-9]+)", sheet.print_title_rows)
        if not match or not 1 <= int(match[1]) <= int(match[2]) < start:
            invalid("print_title_rows")
    if sheet.freeze_panes and sheet[sheet.freeze_panes].row > start:
        invalid("freeze_panes")
    return book


def _print_columns(sheet):
    if not sheet.print_area:
        return 1, sheet.max_column
    areas = PrintArea.from_string(sheet.print_area).ranges
    if len(areas) != 1:
        invalid("print_area")
    area = next(iter(areas))
    bounded_range(str(area))
    if area.min_row != 1:
        invalid("print_area")
    return area.min_col, area.max_col


def _copy_cell(source, target):
    target.value = source.value
    target.data_type = source.data_type
    for attr in ("font", "fill", "border", "alignment", "protection"):
        setattr(target, attr, copy(getattr(source, attr)))
    target.number_format = source.number_format


def _foreground(sheet, cell):
    """Keep simple fills and choose readable text for each business binding."""
    if cell.fill.fill_type is None:
        return "172B4D"
    if cell.fill.fill_type != "solid":
        invalid(cell.coordinate, "STATEMENT_TEMPLATE_MAPPING")
    color = cell.fill.fgColor
    try:
        if color.type == "rgb":
            rgb = color.rgb[-6:]
        elif color.type == "indexed" and color.indexed < 64:
            rgb = COLOR_INDEX[color.indexed][-6:]
        elif color.type == "theme":
            theme = ElementTree.fromstring(sheet.parent.loaded_theme)
            namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
            scheme = theme.find(".//" + namespace + "clrScheme")
            names = [
                "lt1",
                "dk1",
                "lt2",
                "dk2",
                "accent1",
                "accent2",
                "accent3",
                "accent4",
                "accent5",
                "accent6",
                "hlink",
                "folHlink",
            ]
            entry = scheme.find(namespace + names[color.theme])[0]
            rgb = entry.attrib.get("lastClr") or entry.attrib["val"]
        else:
            invalid(cell.coordinate, "STATEMENT_TEMPLATE_MAPPING")
        channels = tuple(int(rgb[i : i + 2], 16) / 255 for i in (0, 2, 4))
        hue, lightness, saturation = rgb_to_hls(*channels)
        tint = color.tint
        lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
        channels = hls_to_rgb(hue, lightness, saturation)
        linear = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in channels]
        luminance = sum(
            v * weight for v, weight in zip(linear, (0.2126, 0.7152, 0.0722), strict=True)
        )
        return "000000" if luminance > 0.179 else "FFFFFF"
    except ValueError, TypeError, IndexError, KeyError, ElementTree.ParseError:
        invalid(cell.coordinate, "STATEMENT_TEMPLATE_MAPPING")


def _column_width(sheet, column):
    for dimension in sheet.column_dimensions.values():
        if dimension.min <= column <= dimension.max:
            return dimension.width or sheet.sheet_format.defaultColWidth or 13
    return sheet.sheet_format.defaultColWidth or 13


def _put(sheet, address, value, number=False, money=False):
    target = sheet[address]
    if value is not None and len(str(value)) > 32767:
        invalid(address, "STATEMENT_TEMPLATE_TEXT_LIMIT")
    target.value = None
    if value is not None:
        if number and len(Decimal(str(value)).as_tuple().digits) <= 15:
            target.value = Decimal(str(value))
            target.number_format = (
                "0.00"
                if money
                else "0"
                if target.value == target.value.to_integral_value()
                else "0.####"
            )
        else:
            target.value = str(value)
            target.data_type = "s"
            target.number_format = "General"
    else:
        target.number_format = "0.00" if money else "General"
    alignment = copy(target.alignment)
    alignment.wrap_text = True
    alignment.shrink_to_fit = False
    alignment.text_rotation = 0
    alignment.horizontal = "right" if number else "left"
    alignment.indent = alignment.relativeIndent = 0
    alignment.readingOrder = 0
    alignment.justifyLastLine = False
    target.alignment = alignment
    # Transparent/white text or a conditional number format cannot hide a business binding.
    target.font = copy(target.font)
    font = copy(target.font)
    font.color, font.sz = _foreground(sheet, target), max(10, font.sz or 10)
    target.font = font


def render_customer_excel(snapshot, content):
    template = snapshot["customer_template"]
    mapping, data = template["mapping"], snapshot["statement"]
    book = validate_mapping(content, mapping)
    sheet = book.active
    print_left, print_right = _print_columns(sheet)
    require_coverage(mapping, data)
    start, end = mapping["detail_start_row"], mapping["detail_end_row"]
    all_lines = data["lines"] + data.get("settlement_lines", [])
    lines = [(i, line) for i, line in enumerate(all_lines) if line["status"] == "included"]
    shift = len(lines) - (end - start + 1)
    original = {
        cell.coordinate: copy(cell)
        for row in sheet
        for cell in row
        if not isinstance(cell, MergedCell)
    }
    dimensions = {row: copy(dim) for row, dim in sheet.row_dimensions.items()}
    merges = [copy(area) for area in sheet.merged_cells.ranges]
    max_row, max_col = sheet.max_row, sheet.max_column
    written = set()
    for area in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(area))
    sheet.delete_rows(1, max_row)
    sheet.row_dimensions.clear()

    def copy_row(source_row, target_row, clear=False):
        if source_row in dimensions:
            dimension = copy(dimensions[source_row])
            dimension.index = target_row
            sheet.row_dimensions[target_row] = dimension
        for col in range(1, max_col + 1):
            old = original.get(f"{get_column_letter(col)}{source_row}")
            if old is not None:
                target = sheet.cell(target_row, col)
                _copy_cell(old, target)
                if clear:
                    target.value = None

    for row in range(1, start):
        copy_row(row, row)
    for row in range(start, start + len(lines)):
        copy_row(start, row, clear=True)
    for row in range(end + 1, max_row + 1):
        copy_row(row, row + shift)
    for area in merges:
        if area.max_row < start:
            sheet.merge_cells(str(area))
        elif area.min_row > end:
            area.shift(row_shift=shift)
            sheet.merge_cells(str(area))
        elif area.min_row == start:
            for offset in range(len(lines)):
                repeated = copy(area)
                repeated.shift(row_shift=offset)
                sheet.merge_cells(str(repeated))
    for ordinal, (index, line) in enumerate(lines, 1):
        source = line["source"]
        values = {
            "ordinal": ordinal,
            "delivery_date": source["delivery_date"],
            "item_code": source.get("item_code"),
            "item_name": source.get("item_name") or source.get("description"),
            "item_name_specification": " / ".join(
                source[key] for key in ("item_name", "specification") if source.get(key)
            )
            or source.get("description"),
            "order_number": source.get("order_number"),
            "delivery_reference": source.get("delivery_id") or line.get("delivery_id"),
            "external_document_number": line.get("external_document_number", ""),
            "specification": source.get("specification"),
            "unit": source.get("canonical_unit") or source.get("source_unit"),
            "quantity": line["quantity"],
            "unit_price": line["unit_price"] if line["unit_price"] is not None else "待定价",
            "amount": data["calculation"]["line_amounts"][index],
            "reason": line["reason"],
        }
        for field, col in mapping["columns"].items():
            value = values[field]
            address = f"{col}{start + ordinal - 1}"
            _put(
                sheet,
                address,
                value,
                number=field in {"quantity", "unit_price", "amount"}
                and value not in {None, "待定价"},
                money=field == "amount",
            )
            written.add(address)
    formal = data.get("status") == "confirmed"
    label = f"{'正式' if formal else '草稿'} v{data['version']}"
    tax_label = {
        "included": "含税不另加税",
        "untaxed_final": "未税不另加税",
        "add_tax": f"未税加税 {data['tax_rate']}%",
    }[data["tax_mode"]]
    label += f" · {tax_label} · CNY"
    if snapshot.get("current") is False:
        label = "历史版本 · " + label
    if snapshot.get("preview"):
        label = "模板预览 · " + label
    if "PRICE_REQUIRED" in data["calculation"]["warnings"]:
        label += " · 金额未确定（缺单价）"
    header_label = label
    if data["tax_mode"] == "add_tax" or data["adjustments"]:
        for field, heading in (("tax_base", "计税基数"), ("tax_amount", "税额")):
            if field not in mapping["cells"]:
                value = data["calculation"][field]
                header_label += f"\n{heading}：{value if value is not None else '待确定'}"
    values = {
        **data["calculation"],
        "customer_name": data["customer_name"],
        "period": f"{data['period_start']} 至 {data['period_end']}",
        "status": header_label,
        "total_label": "未税" if data["tax_mode"] == "untaxed_final" else "含税",
        "adjustments": "\n".join(
            f"{'税前' if a['stage'] == 'before_tax' else '税后'} / {a['kind']} / "
            f"{a['amount']}：{a['reason']}"
            for a in data["adjustments"]
        )
        or "无调整",
    }
    if "status" not in mapping["cells"]:
        values["period"] += "\n" + header_label
    for field, address in mapping["cells"].items():
        old = original[address]
        moved = f"{get_column_letter(old.column)}{old.row + (shift if old.row > end else 0)}"
        _put(
            sheet,
            moved,
            values[field],
            number=field in {"priced_subtotal", "tax_base", "tax_amount", "total"},
            money=True,
        )
        written.add(moved)
    # Wrap business text using actual column/merge widths without discarding text.
    for row in sheet:
        if not any(cell.coordinate in written for cell in row):
            continue
        height = sheet.row_dimensions[row[0].row].height or 20
        for cell in row:
            if cell.coordinate not in written or cell.value is None or not cell.alignment.wrap_text:
                continue
            area = next(
                (area for area in sheet.merged_cells.ranges if cell.coordinate in area), None
            )
            cols = range(area.min_col, area.max_col + 1) if area else [cell.column]
            width = max(1, sum(_column_width(sheet, c) for c in cols) - 2)
            font_size = cell.font.sz or 11
            width = max(1, width * 11 / max(11, font_size))
            lines_count = sum(
                max(1, ceil(sum(2 if ord(c) > 255 else 1 for c in part) / width))
                for part in str(cell.value).split("\n")
            )
            height = max(height, lines_count * max(16, font_size * 1.4))
            if height > 409:
                invalid(cell.coordinate, "STATEMENT_TEMPLATE_TEXT_LIMIT")
        sheet.row_dimensions[row[0].row].height = height
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.page_setup.orientation = sheet.page_setup.orientation or "portrait"
    sheet.page_setup.paperSize = sheet.page_setup.paperSize or sheet.PAPERSIZE_A4
    sheet.print_area = (
        f"{get_column_letter(print_left)}1:{get_column_letter(print_right)}{max_row + shift}"
    )
    if sheet.auto_filter.ref:
        left, _, right, _ = range_boundaries(sheet.auto_filter.ref)
        sheet.auto_filter.ref = (
            f"{get_column_letter(left)}{start - 1}:"
            f"{get_column_letter(right)}{start + len(lines) - 1}"
            if lines
            else None
        )
    sheet.print_title_rows = f"1:{start - 1}"
    identity = (
        f"{label} · 模板 v{template['version']} · 导出 {snapshot['id']} · 发送前核对系统当前版本"
    )
    for footer in (sheet.oddFooter, sheet.evenFooter, sheet.firstFooter):
        footer.center.text = ((footer.center.text + "\n") if footer.center.text else "") + identity
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
